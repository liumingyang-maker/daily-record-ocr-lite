from __future__ import annotations

from openpyxl import load_workbook

from lite_app.knowledge.exporter import export_knowledge_history
from tests_lite.knowledge_fixtures import seed_dated_history


def test_knowledge_export_is_chronological_and_business_labeled(tmp_path):
    db_path = tmp_path / "knowledge.sqlite3"
    seed_dated_history(db_path)
    output = export_knowledge_history(db_path, tmp_path / "配方知识库.xlsx")

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["配方历史"]
    assert [cell.value for cell in sheet[1]][:4] == ["客户", "产品", "日期", "配方"]
    assert [sheet.cell(row, 3).value for row in (2, 3)] == ["2026-07-27", "2026-07-28"]
    assert [sheet.cell(row, 6).value for row in (2, 3)] == ["60", "62"]
