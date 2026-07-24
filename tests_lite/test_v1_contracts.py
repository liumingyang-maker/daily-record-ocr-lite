"""v1.0.0 核心正确性与产品安全验收。"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from lite_app.ocr.base import OCRPage, OCRToken


def _token(
    token_id: str,
    text: str,
    bbox: list[float],
    confidence: float = 0.95,
) -> OCRToken:
    return OCRToken(
        id=token_id,
        text=text,
        confidence=confidence,
        polygon=[
            [bbox[0], bbox[1]],
            [bbox[2], bbox[1]],
            [bbox[2], bbox[3]],
            [bbox[0], bbox[3]],
        ],
        bbox=bbox,
        center_x=(bbox[0] + bbox[2]) / 2,
        center_y=(bbox[1] + bbox[3]) / 2,
    )


def test_mock_result_obeys_versioned_record_schema():
    from lite_app.config import PROJECT_ROOT
    from lite_app.contracts import load_record_schema, validate_record_result

    mock_result = json.loads(
        (PROJECT_ROOT / "config" / "mock_result.json").read_text(encoding="utf-8")
    )
    assert load_record_schema()["$id"].endswith("record-v1.schema.json")
    assert validate_record_result(mock_result).fatal == []


def test_missing_pages_is_fatal_schema_error():
    from lite_app.contracts import validate_record_result

    report = validate_record_result({"schema_version": "record-v1", "warnings": []})
    assert report.fatal
    assert report.ready is False


def test_missing_evidence_metadata_is_normalized_to_review_not_failure(tmp_path):
    from lite_app.config import PROJECT_ROOT
    from lite_app.contracts import validate_record_result
    from lite_app.final_result import FinalResultService, project_final_result

    raw = json.loads((PROJECT_ROOT / "config" / "mock_result.json").read_text(encoding="utf-8"))
    field = raw["pages"][0]["product_sections"][0]["formulas"][0]["materials"][0]["amount"]
    del field["evidence_token_ids"]
    report = validate_record_result(raw)
    assert report.fatal == []
    assert report.reviewable

    final = project_final_result("job-reviewable", raw, {"fields": []})
    FinalResultService(tmp_path).replace(final)
    saved_field = FinalResultService(tmp_path).load()["pages"][0]["product_sections"][0][
        "formulas"
    ][0]["materials"][0]["amount"]
    assert saved_field["evidence_token_ids"] == []
    assert saved_field["status"] == "NEED_REVIEW"
    assert saved_field["review_status"] == "NEED_REVIEW"


def test_empty_pages_and_incomplete_upload_coverage_are_fatal():
    from lite_app.contracts import validate_page_coverage, validate_record_result

    result = {"schema_version": "record-v1", "pages": [], "warnings": []}
    assert validate_record_result(result).fatal
    assert validate_page_coverage(result, expected_pages=2)


def test_visual_probe_contract_does_not_disclose_the_answer():
    from lite_app.vision.probe import (
        PROBE_MARKER,
        PROBE_SCHEMA,
        PROBE_SYSTEM_PROMPT,
        PROBE_USER_PROMPT,
        validate_probe_response,
    )

    text_only_contract = json.dumps(
        {
            "system": PROBE_SYSTEM_PROMPT,
            "user": PROBE_USER_PROMPT,
            "schema": PROBE_SCHEMA,
        },
        ensure_ascii=False,
    )
    assert PROBE_MARKER not in text_only_contract
    assert validate_probe_response(json.dumps({"marker": PROBE_MARKER})) is True
    assert validate_probe_response('{"marker": "VISION-0000"}') is False


def test_record_v1_ids_are_rewritten_locally_and_reused_by_projection():
    from lite_app.config import PROJECT_ROOT
    from lite_app.contracts import normalize_legacy_result
    from lite_app.grouping.service import build_business_entities

    raw = json.loads((PROJECT_ROOT / "config" / "mock_result.json").read_text(encoding="utf-8"))
    normalized = normalize_legacy_result(raw, "job-local")
    formula = normalized["pages"][0]["product_sections"][0]["formulas"][0]
    assert formula["formula_id"] == "job-local__page_001__formula_001"
    assert formula["materials"][0]["material_id"] == "material_001"

    entities = build_business_entities("job-local", normalized)
    assert entities.formulas[0].formula_id == formula["formula_id"]


def test_formula_ids_are_unique_across_sections_on_one_page():
    import copy

    from lite_app.config import PROJECT_ROOT
    from lite_app.contracts import normalize_legacy_result

    raw = json.loads((PROJECT_ROOT / "config" / "mock_result.json").read_text(encoding="utf-8"))
    second = copy.deepcopy(raw["pages"][0]["product_sections"][0])
    raw["pages"][0]["product_sections"].append(second)
    normalized = normalize_legacy_result(raw, "job-local")
    formulas = [
        formula
        for section in normalized["pages"][0]["product_sections"]
        for formula in section["formulas"]
    ]
    assert [formula["formula_id"] for formula in formulas] == [
        "job-local__page_001__formula_001",
        "job-local__page_001__formula_002",
    ]
    assert [formula["formula_sequence"] for formula in formulas] == [1, 2]


def test_page_coverage_requires_upload_order_and_projection_reuses_source_id():
    import copy

    from lite_app.config import PROJECT_ROOT
    from lite_app.contracts import normalize_legacy_result, validate_page_coverage
    from lite_app.grouping.service import build_business_entities

    raw = json.loads((PROJECT_ROOT / "config" / "mock_result.json").read_text(encoding="utf-8"))
    second_page = copy.deepcopy(raw["pages"][0])
    raw["pages"][0]["source_image_index"] = 2
    second_page["source_image_index"] = 1
    raw["pages"].append(second_page)
    normalized = normalize_legacy_result(raw, "job-pages")
    assert validate_page_coverage(normalized, expected_pages=2)

    entities = build_business_entities("job-pages", normalized)
    assert [(page.source_image_index, page.page_id) for page in entities.pages] == [
        (2, "page_002"),
        (1, "page_001"),
    ]


def test_v1_error_taxonomy_is_importable():
    from lite_app.exporter import UnresolvedReviewError
    from lite_app.fusion.association import AssociationError
    from lite_app.fusion.engine import FusionError
    from lite_app.layout.geometry import LayoutError
    from lite_app.vision.base import (
        VisionConfigurationError,
        VisionConnectionError,
    )
    from scripts.doctor import InstallationCheckError

    for error_type in (
        AssociationError,
        LayoutError,
        FusionError,
        UnresolvedReviewError,
        VisionConfigurationError,
        VisionConnectionError,
        InstallationCheckError,
    ):
        assert issubclass(error_type, RuntimeError)


def test_job_statuses_are_centrally_defined():
    from lite_app.status import JobStatus

    assert {status.value for status in JobStatus} >= {
        "UPLOADED",
        "PREPROCESSING",
        "OCR_RUNNING",
        "VISION_RUNNING",
        "MATCHING_HISTORY",
        "FUSING",
        "REVIEW_REQUIRED",
        "READY",
        "FAILED",
        "FAILED_SCHEMA",
        "DEGRADED",
    }


def test_windows_installer_uses_temporary_short_drive_for_paddle_wheels():
    from lite_app.config import PROJECT_ROOT

    script = (PROJECT_ROOT / "install" / "install-windows.ps1").read_text(encoding="utf-8-sig")

    assert "subst.exe $ShortDrive $Root" in script
    assert "subst.exe $ShortDrive /D" in script
    assert '$InstallRoot = "$ShortDrive\\"' in script
    assert "finally" in script


def test_windows_updater_reuses_hardened_installer():
    from lite_app.config import PROJECT_ROOT

    script = (PROJECT_ROOT / "install" / "update-windows.ps1").read_text(encoding="utf-8-sig")

    assert '$Installer = Join-Path $Root "install\\install-windows.ps1"' in script
    assert "& $Installer @InstallArguments" in script


def test_environment_overrides_persisted_settings(tmp_path, monkeypatch):
    import lite_app.config as config_module

    original_root = config_module.PROJECT_ROOT
    (tmp_path / "config").mkdir()
    (tmp_path / "data").mkdir()
    for name in ("app.yaml", "recognition.yaml"):
        (tmp_path / "config" / name).write_text(
            (original_root / "config" / name).read_text(encoding="utf-8"),
            encoding="utf-8",
        )
    (tmp_path / "data" / "settings.json").write_text(
        json.dumps(
            {
                "vision": {"provider": "", "model": ""},
                "ocr": {"provider": "paddleocr_v6"},
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(config_module, "PROJECT_ROOT", tmp_path)
    monkeypatch.setenv("VISION_PROVIDER", "mock")
    monkeypatch.setenv("OCR_PROVIDER", "mock")
    config_module.clear_config_cache()
    try:
        assert config_module.get_config().vision["provider"] == "mock"
        recognition = config_module.load_recognition_config()
        assert recognition["vision"]["provider"] == "mock"
        assert recognition["ocr"]["provider"] == "mock"
    finally:
        config_module.clear_config_cache()


@pytest.mark.parametrize("tier", ["tiny", "small", "medium"])
def test_ppocr_v6_tier_names_are_official(tier):
    from lite_app.ocr.paddleocr_v6 import TIER_MODELS

    assert TIER_MODELS[tier]["text_detection_model_name"] == f"PP-OCRv6_{tier}_det"
    assert TIER_MODELS[tier]["text_recognition_model_name"] == f"PP-OCRv6_{tier}_rec"


def test_unknown_paddle_result_shape_raises():
    from lite_app.ocr.paddleocr_v6 import OCRResultParseError, PaddleOCRv6Provider

    provider = PaddleOCRv6Provider()
    with pytest.raises(OCRResultParseError):
        provider._parse_result(object(), 100, 100)


def test_cpu_provider_disables_mkldnn_for_paddle_33_compatibility(monkeypatch):
    from lite_app.ocr.paddleocr_v6 import PaddleOCRv6Provider

    captured = {}

    class FakePaddleOCR:
        def __init__(self, **kwargs):
            captured.update(kwargs)

    monkeypatch.setitem(sys.modules, "paddleocr", SimpleNamespace(PaddleOCR=FakePaddleOCR))
    PaddleOCRv6Provider(device="cpu").load()

    assert captured["enable_mkldnn"] is False


@pytest.mark.asyncio
async def test_failed_ocr_provider_is_reported_as_unavailable():
    from lite_app.ocr.manager import OCRModelManager, OCRUnavailableError

    class BrokenProvider:
        is_loaded = True
        provider_name = "paddleocr_v6"

        def recognize(self, _image_path):
            raise RuntimeError("predictor crashed")

    OCRModelManager.reset()
    manager = OCRModelManager()
    manager.configure({"enabled": True, "provider": "paddleocr_v6"})
    manager._provider = BrokenProvider()
    with pytest.raises(OCRUnavailableError, match="predictor crashed"):
        await manager.recognize_async(Path("unused.png"))

    status = manager.get_status()
    assert status["configured_provider"] == "paddleocr_v6"
    assert status["effective_provider"] == "unavailable"
    assert "predictor crashed" in status["last_error"]
    OCRModelManager.reset()


def test_page_scoping_rewrites_every_token_id():
    from lite_app.ocr.base import scope_token_ids

    page = OCRPage(
        image_index=1,
        width=100,
        height=100,
        tokens=[_token("p1_t001", "0.15", [1, 1, 20, 10])],
        average_confidence=0.95,
        provider="mock",
        model="mock",
        elapsed_ms=1,
    )
    scope_token_ids(page, 2)
    assert page.image_index == 2
    assert [token.id for token in page.tokens] == ["p2_t001"]


def test_center_distance_is_independent_of_iou():
    from lite_app.fusion.association import FieldEvidence, associate_field

    page = OCRPage(
        image_index=1,
        width=1000,
        height=1000,
        tokens=[_token("p1_t001", "0.15", [120, 120, 160, 145])],
        average_confidence=0.95,
        provider="mock",
        model="mock",
        elapsed_ms=1,
    )
    candidates = associate_field(
        FieldEvidence(
            field_id="amount",
            field_type="amount",
            source_image_index=1,
            field_bbox=[0.158, 0.132, 0.198, 0.157],
            record_bbox=None,
            evidence_token_ids=[],
            vlm_value="0.5",
        ),
        [page],
        None,
    )
    assert candidates
    assert candidates[0].method == "center_distance"
    assert candidates[0].value == "0.15"


def test_multitoken_decimal_candidate_preserves_all_evidence():
    from lite_app.fusion.association import FieldEvidence, associate_field

    page = OCRPage(
        image_index=1,
        width=1000,
        height=1000,
        tokens=[
            _token("p1_t001", "0", [100, 100, 120, 130]),
            _token("p1_t002", ".", [121, 100, 128, 130]),
            _token("p1_t003", "15", [129, 100, 160, 130]),
        ],
        average_confidence=0.95,
        provider="mock",
        model="mock",
        elapsed_ms=1,
    )
    candidates = associate_field(
        FieldEvidence(
            field_id="amount",
            field_type="amount",
            source_image_index=1,
            field_bbox=[0.10, 0.10, 0.16, 0.13],
            record_bbox=None,
            evidence_token_ids=["p1_t001", "p1_t002", "p1_t003"],
            vlm_value="0.5",
        ),
        [page],
        None,
    )
    assert candidates[0].value == "0.15"
    assert candidates[0].token_ids == ["p1_t001", "p1_t002", "p1_t003"]


def test_layout_fallback_selects_material_pair_within_matching_record():
    from lite_app.fusion.association import FieldEvidence, associate_field

    page = OCRPage(
        image_index=1,
        width=1000,
        height=1000,
        tokens=[
            _token("p1_t001", "PA66", [100, 100, 180, 140]),
            _token("p1_t002", "0.15", [100, 170, 180, 210]),
            _token("p1_t003", "GF30", [100, 600, 180, 640]),
            _token("p1_t004", "0.25", [100, 670, 180, 710]),
        ],
        average_confidence=0.95,
        provider="paddleocr_v6",
        model="PP-OCRv6_medium",
        elapsed_ms=1,
    )
    layout = {
        "pairs": {
            "1": [
                {
                    "name_token_ids": ["p1_t001"],
                    "amount_token_ids": ["p1_t002"],
                    "name": "PA66",
                    "score": 0.9,
                },
                {
                    "name_token_ids": ["p1_t003"],
                    "amount_token_ids": ["p1_t004"],
                    "name": "GF30",
                    "score": 0.9,
                },
            ]
        },
        "records": {
            "1": [
                {"bbox": [50, 50, 400, 300]},
                {"bbox": [50, 550, 400, 800]},
            ]
        },
    }
    candidates = associate_field(
        FieldEvidence(
            field_id="amount-2",
            field_type="amount",
            source_image_index=1,
            field_bbox=None,
            record_bbox=[0.05, 0.55, 0.4, 0.8],
            evidence_token_ids=[],
            vlm_value="0.5",
            anchor_value="GF30",
        ),
        [page],
        layout,
    )

    assert len(candidates) == 1
    assert candidates[0].value == "0.25"
    assert candidates[0].token_ids == ["p1_t004"]


def test_record_boundary_mismatch_is_explicit_review_signal():
    from lite_app.fusion.association import record_boundary_mismatch

    page = OCRPage(
        image_index=1,
        width=1000,
        height=1000,
        tokens=[],
        average_confidence=0,
        provider="paddleocr_v6",
        model="PP-OCRv6_medium",
        elapsed_ms=1,
    )
    assert record_boundary_mismatch(
        1,
        [0.7, 0.7, 0.9, 0.9],
        [page],
        {"records": {"1": [{"bbox": [50, 50, 400, 300]}]}},
    )


def test_final_result_manual_update_is_excel_source(tmp_path):
    from lite_app.final_result import FinalResultService, project_final_result
    from lite_app.fusion.engine import CONFLICT
    from lite_app.grouping.exporter import export_grouped_excel
    from lite_app.grouping.service import build_business_entities

    structured = {
        "schema_version": "record-v1",
        "pages": [
            {
                "page_id": "page_001",
                "source_image_index": 1,
                "company": {
                    "raw_value": "测试公司",
                    "standard_value": "测试公司",
                    "confidence": 0.9,
                    "bbox": None,
                    "evidence_token_ids": [],
                    "match_source": "vlm",
                    "review_status": "AUTO_ACCEPT",
                },
                "product_sections": [
                    {
                        "section_id": "section_001",
                        "product_or_series": {
                            "value": "PA66",
                            "confidence": 0.9,
                            "evidence_token_ids": [],
                            "bbox": None,
                        },
                        "product_type": "product",
                        "section_bbox": None,
                        "formulas": [
                            {
                                "formula_id": "job__page_001__formula_001",
                                "formula_no": "①",
                                "formula_sequence": 1,
                                "record_date": {
                                    "value": "",
                                    "confidence": 0.0,
                                    "evidence_token_ids": [],
                                    "bbox": None,
                                },
                                "record_bbox": None,
                                "materials": [
                                    {
                                        "material_id": "material_001",
                                        "name": {
                                            "value": "EBS",
                                            "confidence": 0.9,
                                            "evidence_token_ids": [],
                                            "bbox": None,
                                        },
                                        "amount": {
                                            "value": "0.5",
                                            "confidence": 0.87,
                                            "evidence_token_ids": ["p1_t001"],
                                            "bbox": [0.1, 0.2, 0.2, 0.3],
                                        },
                                        "unit": {
                                            "value": "",
                                            "confidence": 0.0,
                                            "evidence_token_ids": [],
                                            "bbox": None,
                                        },
                                        "warnings": [],
                                    }
                                ],
                                "process_parameters": [],
                                "notes": {
                                    "value": "",
                                    "confidence": 0.0,
                                    "evidence_token_ids": [],
                                    "bbox": None,
                                },
                                "warnings": [],
                                "confidence": 0.8,
                            }
                        ],
                        "warnings": [],
                    }
                ],
                "warnings": [],
            }
        ],
        "warnings": [],
    }
    amount_id = "job__page_001__formula_001__material_001__amount"
    fusion = {
        "fields": [
            {
                "field_id": amount_id,
                "field_type": "amount",
                "final_value": "0.15",
                "final_confidence": 0.91,
                "final_source": "conflict",
                "status": CONFLICT,
                "reasons": ["OCR=0.15, VLM=0.5"],
                "candidates": [
                    {"value": "0.5", "source": "vlm", "confidence": 0.87},
                    {"value": "0.15", "source": "ocr_base", "confidence": 0.91},
                ],
                "bbox": [0.1, 0.2, 0.2, 0.3],
                "source_image_index": 1,
            }
        ]
    }
    final = project_final_result("job", structured, fusion)
    review_dir = tmp_path / "review"
    review_dir.mkdir()
    service = FinalResultService(tmp_path)
    service.save(final)
    service.update_field(amount_id, "0.15", "manual")

    updated = service.load()
    assert (
        updated["pages"][0]["product_sections"][0]["formulas"][0]["materials"][0]["amount"]["value"]
        == "0.15"
    )
    assert (
        structured["pages"][0]["product_sections"][0]["formulas"][0]["materials"][0]["amount"][
            "value"
        ]
        == "0.5"
    )

    entities = build_business_entities("job", updated)
    output = tmp_path / "recognized.xlsx"
    export_grouped_excel(entities, output)
    workbook = load_workbook(output)
    assert workbook["配方明细"].cell(2, 7).value == "0.15"
    workbook.close()


def test_formula_merge_and_split_keep_final_fusion_and_entity_ids_aligned(tmp_path):
    import copy

    from lite_app.config import PROJECT_ROOT
    from lite_app.contracts import normalize_legacy_result
    from lite_app.final_result import FinalResultService, project_final_result
    from lite_app.grouping.review import GroupingReviewService
    from lite_app.grouping.storage import load_business_entities
    from lite_app.knowledge.database import KnowledgeDB
    from lite_app.storage import read_json_required, write_json_atomic

    raw = json.loads((PROJECT_ROOT / "config" / "mock_result.json").read_text(encoding="utf-8"))
    formulas = raw["pages"][0]["product_sections"][0]["formulas"]
    second = copy.deepcopy(formulas[0])
    second["formula_no"] = "②"
    second["materials"][0]["amount"]["value"] = "0.25"
    formulas.append(second)
    structured = normalize_legacy_result(raw, "job-structure")
    final = project_final_result("job-structure", structured, {"fields": []})
    final_service = FinalResultService(tmp_path)
    final_service.replace(final)

    first_id, second_id = [
        formula["formula_id"]
        for formula in final_service.load()["pages"][0]["product_sections"][0]["formulas"]
    ]
    first_amount = f"{first_id}__material_001__amount"
    second_amount = f"{second_id}__material_001__amount"
    final_service.update_field(second_amount, "0.25", "manual")
    database_path = tmp_path / "knowledge.sqlite3"
    database = KnowledgeDB(database_path)
    database.initialize()
    database.add_correction(
        "job-structure",
        second_amount,
        "amount",
        "0.5",
        "0.25",
    )
    database.close()
    write_json_atomic(
        tmp_path / "fusion" / "result.json",
        {
            "fields": [
                {"field_id": first_amount, "final_value": "0.15"},
                {"field_id": second_amount, "final_value": "0.25"},
            ]
        },
    )
    write_json_atomic(
        tmp_path / "review" / "group_overrides.json",
        {
            "schema_version": "group-overrides-v1",
            "formula_groups": {
                first_id: {"company_id": "company-a"},
                second_id: {"company_id": "company-b"},
            },
            "company_merges": {},
            "product_merges": {},
        },
    )

    review = GroupingReviewService(
        tmp_path,
        "job-structure",
        knowledge_db_path=database_path,
    )
    merged_id = review.merge_formulas([first_id, second_id])
    merged = final_service.load()["pages"][0]["product_sections"][0]["formulas"][0]
    assert merged_id == first_id
    assert [item["material_id"] for item in merged["materials"]] == [
        "material_001",
        "material_002",
    ]
    assert [item["amount"]["field_id"] for item in merged["materials"]] == [
        f"{first_id}__material_001__amount",
        f"{first_id}__material_002__amount",
    ]

    new_id = review.split_formula(first_id, 1)
    split_formulas = final_service.load()["pages"][0]["product_sections"][0]["formulas"]
    assert [formula["materials"][0]["material_id"] for formula in split_formulas] == [
        "material_001",
        "material_001",
    ]
    assert split_formulas[1]["materials"][0]["amount"]["field_id"] == (
        f"{new_id}__material_001__amount"
    )

    entities = load_business_entities(tmp_path)
    assert entities is not None
    assert [formula.materials[0].field_id for formula in entities.formulas] == [
        f"{first_id}__material_001",
        f"{new_id}__material_001",
    ]
    fusion_ids = {
        item["field_id"]
        for item in read_json_required(tmp_path / "fusion" / "result.json")["fields"]
    }
    assert fusion_ids == {
        f"{first_id}__material_001__amount",
        f"{new_id}__material_001__amount",
    }
    groups = read_json_required(tmp_path / "review" / "group_overrides.json")["formula_groups"]
    assert set(groups) == {first_id, new_id}
    assert groups[new_id] == groups[first_id]
    correction = read_json_required(tmp_path / "review" / "corrections.json")[0]
    assert correction["original_field_id"] == second_amount
    assert correction["field_id"] == f"{new_id}__material_001__amount"
    database = KnowledgeDB(database_path)
    database.initialize()
    stored_correction = database.get_corrections_for_job("job-structure")[0]
    database.close()
    assert stored_correction["original_field_id"] == second_amount
    assert stored_correction["field_id"] == f"{new_id}__material_001__amount"


def test_settings_secret_is_never_returned(tmp_path, monkeypatch):
    from lite_app.settings import SettingsService

    monkeypatch.setenv("DEMO_MODE", "false")
    service = SettingsService(tmp_path)
    assert service.status()["state"] == "SETUP_REQUIRED"
    service.configure_vision(
        {
            "provider": "openai_compatible",
            "base_url": "https://example.invalid/v1",
            "endpoint": "/chat/completions",
            "model": "vision-model",
            "api_key": "super-secret",
        }
    )
    public = service.public_settings()
    assert "super-secret" not in json.dumps(public)
    assert public["vision"]["api_key_configured"] is True
    assert "api_key" not in public["vision"]


def test_cache_blob_survives_source_job_deletion(tmp_path):
    from lite_app.cache import FileRecognitionCache

    cache = FileRecognitionCache(tmp_path / "cache")
    cache.put_json("ocr", "key", {"tokens": [{"id": "p1_t001"}]})
    assert cache.get_json("ocr", "key")["tokens"][0]["id"] == "p1_t001"


@pytest.mark.asyncio
async def test_cross_job_ocr_cache_avoids_second_provider_call_and_rescopes_tokens(
    tmp_path,
):
    from lite_app.cache import FileRecognitionCache
    from lite_app.pipeline_v2 import _recognize_with_cache

    class CountingManager:
        def __init__(self):
            self.calls = 0
            self.loads = 0

        def ensure_loaded(self):
            self.loads += 1

        async def recognize_async(self, _image_path):
            self.calls += 1
            return OCRPage(
                image_index=1,
                width=100,
                height=100,
                tokens=[_token("t001", "0.15", [10, 10, 30, 20])],
                average_confidence=0.95,
                provider="counting",
                model="PP-OCRv6_medium",
                elapsed_ms=1,
            )

    manager = CountingManager()
    cache = FileRecognitionCache(tmp_path / "cache")
    job_a, hit_a = await _recognize_with_cache(
        manager, cache, tmp_path / "same.png", 1, True, "same-config"
    )
    job_b, hit_b = await _recognize_with_cache(
        manager, cache, tmp_path / "same.png", 2, True, "same-config"
    )
    _, hit_changed = await _recognize_with_cache(
        manager, cache, tmp_path / "same.png", 1, True, "changed-config"
    )

    assert manager.calls == 2
    assert manager.loads == 1
    assert hit_a is False
    assert hit_b is True
    assert hit_changed is False
    assert job_a.tokens[0].id == "p1_t001"
    assert job_b.tokens[0].id == "p2_t001"


def test_ocr_cache_key_changes_with_explicit_model_identity(tmp_path):
    from lite_app.pipeline_v2 import _build_ocr_cache_key

    image = tmp_path / "page.png"
    image.write_bytes(b"same pixels")
    common = {
        "provider": "paddleocr_v6",
        "tier": "medium",
        "device": "cpu",
        "minimum_score": 0.45,
        "use_textline_orientation": True,
    }
    first = _build_ocr_cache_key(
        image,
        "auto",
        {"max_side": 2600},
        common,
        {
            "det_model": "PP-OCRv6_medium_det",
            "rec_model": "PP-OCRv6_medium_rec",
        },
    )
    changed = _build_ocr_cache_key(
        image,
        "auto",
        {"max_side": 2600},
        common,
        {
            "det_model": "PP-OCRv6_small_det",
            "rec_model": "PP-OCRv6_medium_rec",
        },
    )
    assert first != changed
