"""缓存模块和知识库导入测试。"""


from lite_app.cache import RecognitionCache, build_cache_key, compute_image_hash
from lite_app.knowledge.database import KnowledgeDB


class TestCache:
    def test_compute_image_hash(self, tmp_path):
        img = tmp_path / "test.jpg"
        img.write_bytes(b"fake image data")
        h = compute_image_hash(img)
        assert len(h) == 64  # SHA256 hex
        # 相同内容相同哈希
        assert compute_image_hash(img) == h

    def test_different_content_different_hash(self, tmp_path):
        img1 = tmp_path / "a.jpg"
        img2 = tmp_path / "b.jpg"
        img1.write_bytes(b"data1")
        img2.write_bytes(b"data2")
        assert compute_image_hash(img1) != compute_image_hash(img2)

    def test_build_cache_key(self):
        key = build_cache_key("abc123", "ocr", "paddleocr_v6", "medium")
        assert "abc123" in key
        assert "ocr" in key
        assert "paddleocr_v6" in key

    def test_cache_set_and_get(self, tmp_path):
        db = KnowledgeDB(tmp_path / "cache_test.sqlite3")
        db.initialize()
        cache = RecognitionCache(db)

        cache.set("key1", "hash1", "ocr", "/path/result.json", model="v6")
        result = cache.get("key1")
        assert result is not None
        assert result["stage"] == "ocr"
        assert result["result_path"] == "/path/result.json"
        db.close()

    def test_cache_invalidation(self, tmp_path):
        db = KnowledgeDB(tmp_path / "cache_test.sqlite3")
        db.initialize()
        cache = RecognitionCache(db)

        cache.set("key1", "hash1", "ocr", "/path/r1.json")
        cache.set("key2", "hash1", "vision", "/path/r2.json")
        cache.set("key3", "hash2", "ocr", "/path/r3.json")

        # 失效 hash1 的所有缓存
        cache.invalidate_for_image("hash1")
        assert cache.get("key1") is None
        assert cache.get("key2") is None
        # hash2 不受影响
        assert cache.get("key3") is not None
        db.close()

    def test_check_ocr_cache(self, tmp_path):
        db = KnowledgeDB(tmp_path / "cache_test.sqlite3")
        db.initialize()
        cache = RecognitionCache(db)

        img = tmp_path / "test.jpg"
        img.write_bytes(b"image data")

        # 未缓存时返回 None
        assert cache.check_ocr_cache(img, "paddleocr_v6", "medium") is None

        # 缓存后返回结果
        image_hash = compute_image_hash(img)
        key = build_cache_key(image_hash, "ocr", "paddleocr_v6", "medium")
        cache.set(key, image_hash, "ocr", "/path/ocr.json", model="medium")
        result = cache.check_ocr_cache(img, "paddleocr_v6", "medium")
        assert result is not None
        db.close()


class TestKnowledgeImport:
    def test_csv_import(self, tmp_path):
        """测试 CSV 导入物料。"""
        csv_file = tmp_path / "materials.csv"
        csv_file.write_text(
            "物料名称,单位,分类,别名\nPA66,kg,尼龙,聚酰胺66;PA-66\nGF30,kg,玻纤,\n",
            encoding="utf-8",
        )

        import csv
        db = KnowledgeDB(tmp_path / "import_test.sqlite3")
        db.initialize()

        # 模拟 CSV 导入逻辑
        with open(csv_file, encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            count = 0
            for row in reader:
                name = (row.get("物料名称") or "").strip()
                if name:
                    unit = (row.get("单位") or "").strip()
                    category = (row.get("分类") or "").strip()
                    mid = db.add_material(name, category, unit)
                    alias_str = (row.get("别名") or "").strip()
                    if alias_str:
                        for alias in alias_str.replace("；", ";").split(";"):
                            if alias.strip():
                                db.add_alias(mid, alias.strip(), alias_type="import", source="csv")
                    count += 1

        assert count == 2
        # 验证 PA66 存在
        result = db.find_material_by_name("PA66")
        assert result is not None
        assert result["standard_name"] == "PA66"
        # 验证别名
        result2 = db.find_material_by_name("聚酰胺66")
        assert result2 is not None
        assert result2["standard_name"] == "PA66"
        db.close()

    def test_xlsx_import(self, tmp_path):
        """测试 Excel 导入物料。"""
        from openpyxl import Workbook
        xlsx_file = tmp_path / "materials.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["物料名称", "单位", "分类", "别名"])
        ws.append(["PA66", "kg", "尼龙", "聚酰胺66"])
        ws.append(["抗氧剂1010", "", "助剂", ""])
        wb.save(str(xlsx_file))


        from openpyxl import load_workbook
        db = KnowledgeDB(tmp_path / "import_test2.sqlite3")
        db.initialize()

        wb2 = load_workbook(str(xlsx_file), read_only=True)
        ws2 = wb2.active
        rows = list(ws2.iter_rows(values_only=True))
        name_col = 0
        unit_col = 1
        category_col = 2
        alias_col = 3

        count = 0
        for row in rows[1:]:
            name = str(row[name_col] or "").strip()
            if not name:
                continue
            unit = str(row[unit_col] or "").strip()
            category = str(row[category_col] or "").strip()
            mid = db.add_material(name, category, unit)
            alias_str = str(row[alias_col] or "").strip()
            if alias_str:
                for alias in alias_str.split(";"):
                    if alias.strip():
                        db.add_alias(mid, alias.strip(), alias_type="import", source="xlsx")
            count += 1
        wb2.close()

        assert count == 2
        result = db.find_material_by_name("抗氧剂1010")
        assert result is not None
        db.close()
