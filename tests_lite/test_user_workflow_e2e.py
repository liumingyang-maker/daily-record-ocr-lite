from __future__ import annotations

import copy
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from lite_app.final_result import FinalResultService, project_final_result
from lite_app.settings import SettingsService
from lite_app.storage import JobStorage


class _IdleQueue:
    def set_handler(self, _handler):
        return None

    async def start(self):
        return None

    async def stop(self):
        return None

    async def submit(self, _job_id):
        return None


def _image_bytes(color: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (80, 60), color).save(buffer, "PNG")
    return buffer.getvalue()


def _two_formula_result(job_id: str) -> dict:
    root = Path(__file__).resolve().parents[1]
    structured = json.loads((root / "config" / "mock_result.json").read_text("utf-8"))
    first_page = structured["pages"][0]
    first_page["company"]["raw_value"] = "联创"
    first_page["company"]["standard_value"] = "联创"
    first_section = first_page["product_sections"][0]
    first_section["product_or_series"]["value"] = "G30A"
    first_formula = first_section["formulas"][0]
    first_formula["formula_id"] = f"{job_id}__page_001__formula_001"
    first_formula["formula_no"] = "配方1"

    second_page = copy.deepcopy(first_page)
    second_page["page_id"] = "page_002"
    second_page["source_image_index"] = 2
    second_section = second_page["product_sections"][0]
    second_section["section_id"] = "section_002"
    second_formula = second_section["formulas"][0]
    second_formula["formula_id"] = f"{job_id}__page_002__formula_001"
    second_formula["formula_no"] = "配方2"
    second_formula["record_date"]["value"] = ""
    structured["pages"].append(second_page)
    return structured


def test_two_image_user_journey_confirms_history_compares_and_exports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    jobs_dir = tmp_path / "jobs"
    data_dir = tmp_path / "data"
    monkeypatch.setenv("JOBS_DIR", str(jobs_dir))
    monkeypatch.setenv("KNOWLEDGE_DB_PATH", str(data_dir / "knowledge.sqlite3"))
    monkeypatch.setenv("DEMO_MODE", "true")

    from lite_app import jobs, main
    from lite_app.config import clear_config_cache

    idle_queue = _IdleQueue()
    monkeypatch.setattr(jobs, "get_task_queue", lambda: idle_queue)
    monkeypatch.setattr(main, "_get_settings", lambda: SettingsService(data_dir))
    clear_config_cache()
    with TestClient(main.app) as client:
        created = client.post(
            "/jobs",
            files=[
                ("files", ("first.png", _image_bytes("white"), "image/png")),
                ("files", ("second.png", _image_bytes("lightgray"), "image/png")),
            ],
            data={"rotation": "0", "rotation_manifest": '["0", "90cw"]'},
            follow_redirects=False,
        )
        assert created.status_code == 303, created.text
        job_id = created.headers["location"].split("/jobs/")[1]
        storage = JobStorage(jobs_dir)
        job = storage.get_job(job_id)
        assert job["rotations"] == ["0", "90cw"]

        job.update(
            {
                "demo_mode": False,
                "status": "REVIEW_REQUIRED",
                "recognition_run_id": "journey-run",
                "final_result_run_id": "journey-run",
                "ocr_engine": {"effective_provider": "paddleocr_v6", "loaded": True},
                "vision_engine": {
                    "provider": "openai_compatible",
                    "model": "qwen3.7-plus",
                    "healthy": True,
                },
            }
        )
        storage.save_job(job)
        final = project_final_result(job_id, _two_formula_result(job_id), {"fields": []})
        final["recognition_run_id"] = "journey-run"
        second = final["pages"][1]["product_sections"][0]["formulas"][0]
        second["record_date"]["status"] = "EMPTY"
        second["materials"][0]["amount"]["status"] = "CONFLICT"
        FinalResultService(storage.get_job_dir(job_id)).replace(final)

        view = client.get(f"/api/jobs/{job_id}/review").json()
        formulas = [formula for group in view["groups"] for formula in group["formulas"]]
        assert formulas[0]["formula_no"] == "配方2"
        assert formulas[0]["collapsed"] is False
        assert next(item for item in formulas if item["formula_no"] == "配方1")["collapsed"] is True
        version = view["version"]
        by_no = {item["formula_no"]: item for item in formulas}

        identity = client.patch(
            f"/api/jobs/{job_id}/review/groups/{view['groups'][0]['id']}",
            json={"version": version, "customer": "联创", "product": "G30A"},
        ).json()
        version = identity["version"]
        for formula_no, record_date in (("配方1", "2026-07-27"), ("配方2", "2026-07-28")):
            saved = client.patch(
                f"/api/jobs/{job_id}/review/formulas/{by_no[formula_no]['id']}",
                json={"version": version, "record_date": record_date},
            ).json()
            version = saved["version"]
        saved = client.patch(
            f"/api/jobs/{job_id}/review/formulas/{by_no['配方2']['id']}/materials/material_001",
            json={"version": version, "amount": "62"},
        ).json()
        version = saved["version"]

        first_id = by_no["配方1"]["id"]
        added_one = client.post(
            f"/api/jobs/{job_id}/review/formulas/{first_id}/materials",
            json={"version": version, "name": "色粉", "amount": "1", "unit": "kg"},
        ).json()
        added_two = client.post(
            f"/api/jobs/{job_id}/review/formulas/{first_id}/materials",
            json={"version": added_one["version"], "name": "助剂", "amount": "2", "unit": "kg"},
        ).json()
        reordered = client.post(
            f"/api/jobs/{job_id}/review/formulas/{first_id}/materials/reorder",
            json={
                "version": added_two["version"],
                "material_ids": [added_two["material_id"], "material_001", added_one["material_id"]],
            },
        ).json()
        deleted = client.request(
            "DELETE",
            f"/api/jobs/{job_id}/review/formulas/{first_id}/materials/{added_one['material_id']}",
            json={"version": reordered["version"]},
        ).json()

        process_one = client.post(
            f"/api/jobs/{job_id}/review/formulas/{first_id}/process",
            json={"version": deleted["version"], "name": "温度", "value": "260", "unit": "℃"},
        ).json()
        process_two = client.post(
            f"/api/jobs/{job_id}/review/formulas/{first_id}/process",
            json={"version": process_one["version"], "name": "转速", "value": "50", "unit": "Hz"},
        ).json()
        edited = client.patch(
            f"/api/jobs/{job_id}/review/formulas/{first_id}/process/{process_one['parameter_id']}",
            json={"version": process_two["version"], "value": "265"},
        ).json()
        deleted_process = client.request(
            "DELETE",
            f"/api/jobs/{job_id}/review/formulas/{first_id}/process/{process_two['parameter_id']}",
            json={"version": edited["version"]},
        ).json()
        version = deleted_process["version"]

        for formula_id in (first_id, by_no["配方2"]["id"]):
            confirmed = client.post(
                f"/api/jobs/{job_id}/review/formulas/{formula_id}/confirm",
                json={"version": version},
            )
            assert confirmed.status_code == 200
            version = confirmed.json()["version"]

        finalized = client.post(f"/api/jobs/{job_id}/finalize")
        assert finalized.status_code == 200
        receipt = finalized.json()["receipt"]
        assert receipt["formula_count"] == 2
        assert storage.get_job(job_id)["status"] == "READY"

        tree = client.get("/api/knowledge/tree?q=G30A").json()
        timeline = tree["customers"][0]["products"][0]["formulas"]
        assert [item["record_date"] for item in timeline] == ["2026-07-27", "2026-07-28"]
        compared = client.get(
            f"/api/knowledge/compare?left={timeline[0]['id']}&right={timeline[1]['id']}"
        ).json()
        assert compared["materials"]["PA66"] == {
            "before": "60",
            "after": "62",
            "unit_before": "",
            "unit_after": "",
        }

        exported = client.post(f"/api/jobs/{job_id}/export")
        assert exported.status_code == 200
        workbook = load_workbook(
            storage.get_job_dir(job_id) / exported.json()["filename"], data_only=True
        )
        values = {
            cell.value
            for row in workbook["配方明细"].iter_rows()
            for cell in row
            if cell.value is not None
        }
        assert "62" in {str(item) for item in values}

        repeated = client.post(f"/api/jobs/{job_id}/finalize")
        assert repeated.status_code == 200
        assert repeated.json()["receipt"] == receipt
        tree = client.get("/api/knowledge/tree?q=G30A").json()
        assert tree["customers"][0]["products"][0]["formula_count"] == 2
    clear_config_cache()
