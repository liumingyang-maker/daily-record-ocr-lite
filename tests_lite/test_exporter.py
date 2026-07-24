"""Excel export tests with FinalResult as the only formal source."""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from lite_app.exporter import ExportError, _resolve_value, export_job
from lite_app.final_result import FinalResultService, project_final_result
from lite_app.readiness import iter_final_fields


class TestResolveValue:
    def test_literal(self):
        assert _resolve_value("literal:固定文字", {}, {}, None, 0, 0) == "固定文字"

    def test_index(self):
        assert _resolve_value("$index", {}, {}, None, 3, 0) == 3

    def test_parent_index(self):
        assert _resolve_value("$parent_index", {}, {}, None, 0, 2) == 2

    def test_root_field(self):
        root = {"page_heading": "标题"}
        assert _resolve_value("$root.page_heading", {}, root, None, 0, 0) == "标题"

    def test_parent_field(self):
        parent = {"title": "配方X"}
        assert _resolve_value("$parent.title", {}, {}, parent, 0, 0) == "配方X"

    def test_normal_field(self):
        item = {"name": "PA66", "amount": "60"}
        assert _resolve_value("name", item, {}, None, 0, 0) == "PA66"

    def test_list_field_as_json(self):
        item = {"warnings": ["a", "b"]}
        result = _resolve_value("warnings", item, {}, None, 0, 0)
        assert result == '["a", "b"]'


@pytest.fixture
def strict_result() -> dict:
    root = Path(__file__).resolve().parents[1]
    result = json.loads((root / "config" / "mock_result.json").read_text("utf-8"))
    page = result["pages"][0]
    page["company"]["raw_value"] = "测试公司"
    page["company"]["standard_value"] = "测试公司"
    section = page["product_sections"][0]
    section["product_or_series"]["value"] = "测试系列"
    formula = section["formulas"][0]
    formula["record_date"]["value"] = "24.7.10"
    first = formula["materials"][0]
    first["name"]["value"] = "PA66"
    first["amount"]["value"] = "60"
    first["unit"]["value"] = "kg"
    second = copy.deepcopy(first)
    second["material_id"] = "material_002"
    second["name"]["value"] = "GF30"
    second["amount"]["value"] = "30"
    second["unit"]["value"] = ""
    formula["materials"].append(second)
    formula["process_parameters"] = [
        {
            "parameter_id": "parameter_001",
            "name": {
                "value": "转速",
                "confidence": 0.95,
                "evidence_token_ids": [],
                "bbox": None,
            },
            "value": {
                "value": "50",
                "confidence": 0.95,
                "evidence_token_ids": [],
                "bbox": None,
            },
            "unit": {
                "value": "Hz",
                "confidence": 0.95,
                "evidence_token_ids": [],
                "bbox": None,
            },
            "warnings": [],
        }
    ]
    formula["notes"]["value"] = "备注内容"
    return result


class TestExportJob:
    @pytest.fixture
    def job_with_result(self, storage, strict_result):
        job = storage.create_job()
        strict_result["pages"][0]["product_sections"][0]["formulas"][0][
            "formula_id"
        ] = f"{job['id']}__page_001__formula_001"
        final = project_final_result(job["id"], strict_result, {"fields": []})
        final["recognition_run_id"] = "export-test-run"
        for _, field in iter_final_fields(final):
            field["status"] = "MANUAL_CONFIRMED"
            field["review_status"] = "MANUAL_CONFIRMED"
        FinalResultService(storage.get_job_dir(job["id"])).replace(final)
        job["status"] = "READY"
        job["demo_mode"] = False
        job["images"] = [{"source": "source/source_01_test.jpg"}]
        job["recognition_run_id"] = "export-test-run"
        job["final_result_run_id"] = "export-test-run"
        job["ocr_engine"] = {
            "effective_provider": "paddleocr_v6",
            "loaded": True,
        }
        job["vision_engine"] = {
            "provider": "openai_compatible",
            "model": "vision-model",
            "healthy": True,
        }
        storage.save_job(job)
        return job

    def test_export_creates_file(self, job_with_result, storage):
        filename = export_job(job_with_result["id"], storage)
        assert filename.endswith(".xlsx")
        assert (storage.get_job_dir(job_with_result["id"]) / filename).exists()

    def test_export_creates_five_sheets(self, job_with_result, storage):
        filename = export_job(job_with_result["id"], storage)
        workbook = load_workbook(storage.get_job_dir(job_with_result["id"]) / filename)
        assert workbook.sheetnames == [
            "记录汇总",
            "配方明细",
            "工艺参数",
            "识别审查",
            "修正日志",
        ]

    def test_export_summary(self, job_with_result, storage):
        filename = export_job(job_with_result["id"], storage)
        worksheet = load_workbook(
            storage.get_job_dir(job_with_result["id"]) / filename
        )["记录汇总"]
        assert worksheet["A1"].value == "公司"
        assert worksheet["A2"].value == "测试公司"
        assert worksheet["B2"].value == "测试系列"
        assert worksheet["D2"].value == "24.7.10"

    def test_export_materials_expand(self, job_with_result, storage):
        filename = export_job(job_with_result["id"], storage)
        worksheet = load_workbook(
            storage.get_job_dir(job_with_result["id"]) / filename
        )["配方明细"]
        assert worksheet["F1"].value == "原料名称"
        assert worksheet["G1"].value == "数量"
        assert worksheet["F2"].value == "PA66"
        assert worksheet["G2"].value == "60"
        assert worksheet["F3"].value == "GF30"
        assert worksheet["G3"].value == "30"

    def test_export_process_expand(self, job_with_result, storage):
        filename = export_job(job_with_result["id"], storage)
        worksheet = load_workbook(
            storage.get_job_dir(job_with_result["id"]) / filename
        )["工艺参数"]
        assert worksheet["F2"].value == "转速"
        assert worksheet["G2"].value == "50"
        assert worksheet["H2"].value == "Hz"

    def test_export_escapes_untrusted_excel_formulas(self, job_with_result, storage):
        job_dir = storage.get_job_dir(job_with_result["id"])
        service = FinalResultService(job_dir)
        final = service.load()
        field = final["pages"][0]["product_sections"][0]["formulas"][0][
            "materials"
        ][0]["name"]
        field["value"] = '=HYPERLINK("https://evil.example","click")'
        field["status"] = "MANUAL_CONFIRMED"
        service.replace(final)

        filename = export_job(job_with_result["id"], storage)
        cell = load_workbook(job_dir / filename)["配方明细"]["F2"]
        assert cell.data_type == "s"
        assert cell.value.startswith("'=HYPERLINK")

    def test_export_updates_job_status(self, job_with_result, storage):
        export_job(job_with_result["id"], storage)
        job = storage.get_job(job_with_result["id"])
        assert job["status"] == "EXPORTED"
        assert job["export_file"] is not None

    def test_failed_job_cannot_export_stale_final_result(
        self, job_with_result, storage
    ):
        job_with_result["status"] = "FAILED_SCHEMA"
        job_with_result["validation_errors"] = [{"severity": "fatal"}]
        storage.save_job(job_with_result)

        with pytest.raises(ExportError, match="READY"):
            export_job(job_with_result["id"], storage)
        assert storage.get_job(job_with_result["id"])["status"] == "FAILED_SCHEMA"

    def test_company_review_status_blocks_formal_export(
        self, job_with_result, storage
    ):
        service = FinalResultService(storage.get_job_dir(job_with_result["id"]))
        final = service.load()
        final["pages"][0]["company"]["review_status"] = "NEED_REVIEW"
        service.replace(final)

        with pytest.raises(ExportError, match="未解决字段"):
            export_job(job_with_result["id"], storage)

    def test_export_no_final_result_raises(self, storage):
        job = storage.create_job()
        with pytest.raises(ExportError, match="final_result"):
            export_job(job["id"], storage)

    def test_compat_result_cannot_override_final_result(
        self, job_with_result, storage
    ):
        job_dir = storage.get_job_dir(job_with_result["id"])
        compatibility = json.loads((job_dir / "result.json").read_text("utf-8"))
        compatibility["pages"][0]["product_sections"][0]["formulas"][0][
            "materials"
        ][0]["amount"]["value"] = "999"
        (job_dir / "result.json").write_text(
            json.dumps(compatibility, ensure_ascii=False), encoding="utf-8"
        )

        filename = export_job(job_with_result["id"], storage)
        worksheet = load_workbook(job_dir / filename)["配方明细"]
        assert worksheet["G2"].value == "60"
        assert worksheet["G2"].value != "999"

    def test_output_in_job_dir(self, job_with_result, storage):
        filename = export_job(job_with_result["id"], storage)
        output_path = storage.get_job_dir(job_with_result["id"]) / filename
        assert output_path.exists()
        assert output_path.parent == storage.get_job_dir(job_with_result["id"]) / "export"
