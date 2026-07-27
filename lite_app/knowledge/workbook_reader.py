"""Read XLSX and legacy XLS files into one immutable cell contract."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

CellValue = str | int | float | bool | None
WorkbookFormat = Literal["xlsx", "xls"]


class WorkbookReadError(RuntimeError):
    """A supported workbook could not be read safely."""


@dataclass(frozen=True)
class CellData:
    """One non-empty workbook cell."""

    row: int
    column: int
    value: CellValue
    is_date: bool = False
    number_format: str = ""


@dataclass(frozen=True)
class SheetData:
    """Sparse immutable sheet data with one-based cell lookup."""

    name: str
    max_row: int
    max_column: int
    cells: tuple[CellData, ...]
    hidden: bool = False

    def cell_at(self, row: int, column: int) -> CellData | None:
        return next(
            (
                cell
                for cell in self.cells
                if cell.row == row and cell.column == column
            ),
            None,
        )

    def value_at(self, row: int, column: int) -> CellValue:
        cell = self.cell_at(row, column)
        return cell.value if cell is not None else None

    def row_cells(self, row: int) -> tuple[CellData, ...]:
        return tuple(
            sorted(
                (cell for cell in self.cells if cell.row == row),
                key=lambda cell: cell.column,
            )
        )


@dataclass(frozen=True)
class WorkbookData:
    """Read-only workbook representation used by the legacy extractor."""

    path: Path
    format: WorkbookFormat
    sheets: tuple[SheetData, ...]


def read_workbook(path: Path) -> WorkbookData:
    """Read one supported workbook without modifying its bytes."""
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix == ".xlsx":
        return _read_xlsx(source)
    if suffix == ".xls":
        return _read_xls(source)
    raise ValueError(f"不支持的工作簿格式: {source.suffix or '<none>'}")


def _read_xlsx(path: Path) -> WorkbookData:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise WorkbookReadError(f"无法读取 XLSX: {path.name}: {exc}") from exc

    sheets: list[SheetData] = []
    try:
        for source_sheet in workbook.worksheets:
            cells: list[CellData] = []
            for row in source_sheet.iter_rows():
                for source_cell in row:
                    if source_cell.value is None:
                        continue
                    value, is_date = _normalize_openpyxl_value(source_cell)
                    cells.append(
                        CellData(
                            row=int(source_cell.row),
                            column=int(source_cell.column),
                            value=value,
                            is_date=is_date,
                            number_format=str(source_cell.number_format or ""),
                        )
                    )
            sheets.append(
                SheetData(
                    name=str(source_sheet.title),
                    max_row=int(source_sheet.max_row or 0),
                    max_column=int(source_sheet.max_column or 0),
                    cells=tuple(cells),
                    hidden=source_sheet.sheet_state != "visible",
                )
            )
    finally:
        workbook.close()

    return WorkbookData(path=path.resolve(), format="xlsx", sheets=tuple(sheets))


def _read_xls(path: Path) -> WorkbookData:
    try:
        import xlrd

        workbook = xlrd.open_workbook(path, on_demand=True)
    except Exception as exc:
        raise WorkbookReadError(f"无法读取 XLS: {path.name}: {exc}") from exc

    sheets: list[SheetData] = []
    try:
        for source_sheet in workbook.sheets():
            cells: list[CellData] = []
            for row_index in range(source_sheet.nrows):
                for column_index in range(source_sheet.ncols):
                    source_cell = source_sheet.cell(row_index, column_index)
                    value, is_date = _normalize_xlrd_value(
                        source_cell,
                        workbook.datemode,
                    )
                    if value is None or value == "":
                        continue
                    cells.append(
                        CellData(
                            row=row_index + 1,
                            column=column_index + 1,
                            value=value,
                            is_date=is_date,
                        )
                    )
            sheets.append(
                SheetData(
                    name=str(source_sheet.name),
                    max_row=int(source_sheet.nrows),
                    max_column=int(source_sheet.ncols),
                    cells=tuple(cells),
                    hidden=bool(getattr(source_sheet, "visibility", 0)),
                )
            )
    finally:
        workbook.release_resources()

    return WorkbookData(path=path.resolve(), format="xls", sheets=tuple(sheets))


def _normalize_openpyxl_value(cell: Any) -> tuple[CellValue, bool]:
    value = cell.value
    if bool(getattr(cell, "is_date", False)) and isinstance(
        value,
        (date, datetime),
    ):
        return _iso_date_value(value), True
    return _normalize_scalar(value), False


def _normalize_xlrd_value(cell: Any, datemode: int) -> tuple[CellValue, bool]:
    import xlrd

    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return None, False
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate_as_datetime(cell.value, datemode)
        return _iso_date_value(value), True
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value), False
    return _normalize_scalar(cell.value), False


def _iso_date_value(value: date | datetime) -> str:
    if isinstance(value, datetime) and value.time().isoformat() != "00:00:00":
        return value.isoformat()
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()


def _normalize_scalar(value: Any) -> CellValue:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    return str(value)
