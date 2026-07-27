"""Render private workbook cell ranges as tight and contextual PNG evidence."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

from .import_models import FormulaCandidate
from .import_records import formula_fingerprint
from .workbook_reader import SheetData

BACKGROUND = "#ffffff"
GRID = "#cbd5e1"
HEADER = "#e2e8f0"
TEXT = "#172033"
ROW_HEIGHT = 38
ROW_HEADER_WIDTH = 46
TOP_HEADER_HEIGHT = 32


@dataclass(frozen=True)
class RenderedEvidence:
    tight_path: Path
    context_path: Path
    tight_range: str
    context_range: str


def render_formula_evidence(
    sheet: SheetData,
    formula: FormulaCandidate,
    output_dir: Path,
    source_id: str,
) -> RenderedEvidence:
    """Render exact formula cells and one-row/column surrounding context."""
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(
        f"{formula_fingerprint(formula)}\0{source_id}".encode("utf-8")
    ).hexdigest()
    tight_path = output / f"{digest[:24]}_tight.png"
    context_path = output / f"{digest[:24]}_context.png"

    span = formula.source_span
    tight_bounds = (
        span.start_row,
        span.end_row,
        span.start_column,
        span.end_column,
    )
    context_bounds = (
        max(1, span.start_row - 1),
        min(sheet.max_row, span.end_row + 1),
        max(1, span.start_column - 1),
        min(sheet.max_column, span.end_column + 1),
    )
    _render_range(sheet, tight_bounds, tight_path)
    _render_range(sheet, context_bounds, context_path)
    return RenderedEvidence(
        tight_path=tight_path,
        context_path=context_path,
        tight_range=_range_label(*tight_bounds),
        context_range=_range_label(*context_bounds),
    )


def _render_range(
    sheet: SheetData,
    bounds: tuple[int, int, int, int],
    output_path: Path,
) -> None:
    start_row, end_row, start_column, end_column = bounds
    font = _load_font()
    columns = list(range(start_column, end_column + 1))
    widths = [
        _column_width(sheet, column, start_row, end_row, font)
        for column in columns
    ]
    width = ROW_HEADER_WIDTH + sum(widths)
    height = TOP_HEADER_HEIGHT + (end_row - start_row + 1) * ROW_HEIGHT
    image = Image.new("RGB", (max(width, 120), max(height, 60)), BACKGROUND)
    draw = ImageDraw.Draw(image)

    x = ROW_HEADER_WIDTH
    for column, column_width in zip(columns, widths, strict=True):
        draw.rectangle(
            (x, 0, x + column_width, TOP_HEADER_HEIGHT),
            fill=HEADER,
            outline=GRID,
        )
        draw.text(
            (x + 8, 7),
            get_column_letter(column),
            fill=TEXT,
            font=font,
        )
        x += column_width

    for row_offset, row in enumerate(range(start_row, end_row + 1)):
        y = TOP_HEADER_HEIGHT + row_offset * ROW_HEIGHT
        draw.rectangle(
            (0, y, ROW_HEADER_WIDTH, y + ROW_HEIGHT),
            fill=HEADER,
            outline=GRID,
        )
        draw.text((8, y + 9), str(row), fill=TEXT, font=font)
        x = ROW_HEADER_WIDTH
        for column, column_width in zip(columns, widths, strict=True):
            draw.rectangle(
                (x, y, x + column_width, y + ROW_HEIGHT),
                fill=BACKGROUND,
                outline=GRID,
            )
            value = sheet.value_at(row, column)
            text = "" if value is None else str(value)
            draw.text((x + 8, y + 9), text, fill=TEXT, font=font)
            x += column_width

    image.save(output_path, "PNG", optimize=True)


def _column_width(
    sheet: SheetData,
    column: int,
    start_row: int,
    end_row: int,
    font: ImageFont.ImageFont | ImageFont.FreeTypeFont,
) -> int:
    values = [
        "" if (value := sheet.value_at(row, column)) is None else str(value)
        for row in range(start_row, end_row + 1)
    ]
    boxes = [font.getbbox(value or " ") for value in values]
    text_width = max((box[2] - box[0] for box in boxes), default=40)
    return max(80, min(text_width + 20, 260))


def _load_font() -> ImageFont.ImageFont | ImageFont.FreeTypeFont:
    candidates = (
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    )
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), 17)
    try:
        return ImageFont.truetype("DejaVuSans.ttf", 17)
    except OSError:
        return ImageFont.load_default()


def _range_label(
    start_row: int,
    end_row: int,
    start_column: int,
    end_column: int,
) -> str:
    return (
        f"{get_column_letter(start_column)}{start_row}:"
        f"{get_column_letter(end_column)}{end_row}"
    )
