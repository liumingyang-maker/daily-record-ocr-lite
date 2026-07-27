"""Excel export for chronological customer/product formula history."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..excel_safety import safe_excel_value
from .history import KnowledgeHistory

HEADERS = ["客户", "产品", "日期", "配方", "材料", "数量", "单位", "工艺", "确认时间"]


def export_knowledge_history(db_path: Path, output_path: Path) -> Path:
    history = KnowledgeHistory(db_path)
    try:
        formulas = _all_formula_ids(history)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "配方历史"
        for column, header in enumerate(HEADERS, 1):
            cell = sheet.cell(1, column, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        row_number = 2
        for formula_id in formulas:
            detail = history.formula_detail(formula_id)
            process = "；".join(
                f"{item['name']} {item['value']}{item['unit'] or ''}"
                for item in detail["process"]
            )
            materials = detail["materials"] or [{"name": "", "amount": "", "unit": ""}]
            for material in materials:
                values = [
                    detail["customer"],
                    detail["product"],
                    detail["record_date"] or "历史数据未记录日期",
                    detail["formula_no"],
                    material["name"],
                    material["amount"],
                    material["unit"],
                    process,
                    detail["confirmed_at"],
                ]
                for column, value in enumerate(values, 1):
                    sheet.cell(row_number, column, safe_excel_value(value))
                row_number += 1
        sheet.freeze_panes = "A2"
        for column, width in enumerate((18, 18, 14, 14, 22, 12, 10, 30, 24), 1):
            sheet.column_dimensions[chr(64 + column)].width = width
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path
    finally:
        history.close()


def _all_formula_ids(history: KnowledgeHistory) -> list[int]:
    rows = history.database._get_conn().execute(
        """
        SELECT id FROM formulas
        ORDER BY COALESCE(record_date, ''), COALESCE(confirmed_at, ''), id
        """
    ).fetchall()
    return [int(row["id"]) for row in rows]
