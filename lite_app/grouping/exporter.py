"""分组感知的 Excel 导出：含公司/产品/配方编号列。"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..excel_safety import safe_excel_value
from .models import BusinessEntities

logger = logging.getLogger(__name__)

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
_CONFLICT_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def export_grouped_excel(entities: BusinessEntities, output_path: Path) -> None:
    """导出分组感知的 5-Sheet Excel。"""
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    _write_summary_sheet(wb, entities)
    _write_materials_sheet(wb, entities)
    _write_process_sheet(wb, entities)
    _write_review_sheet(wb, entities)
    _write_correction_sheet_placeholder(wb)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    logger.info("分组 Excel 导出完成: %s", output_path)


def _write_header(ws, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _write_value(ws, row: int, column: int, value) -> None:
    ws.cell(row=row, column=column, value=safe_excel_value(value))


def _get_company_name(entities: BusinessEntities, company_id: str) -> str:
    for cg in entities.company_groups:
        if cg.company_id == company_id:
            return cg.display_name
    return ""


def _get_product_name(entities: BusinessEntities, product_id: str) -> str:
    for pg in entities.product_groups:
        if pg.product_id == product_id:
            return pg.display_name
    return ""


def _write_summary_sheet(wb: Workbook, entities: BusinessEntities) -> None:
    """记录汇总：每条配方一行。"""
    ws = wb.create_sheet("记录汇总")
    headers = ["公司", "产品/系列", "配方编号", "日期", "原料数", "工艺参数数", "来源图片", "状态"]
    _write_header(ws, headers)

    row = 2
    for f in entities.formulas:
        company = _get_company_name(entities, f.company_id)
        product = _get_product_name(entities, f.product_id)
        no = f.formula_no_raw or f"未编号{f.formula_sequence}"
        _write_value(ws, row, 1, company)
        _write_value(ws, row, 2, product)
        _write_value(ws, row, 3, no)
        _write_value(ws, row, 4, f.record_date.raw_value)
        _write_value(ws, row, 5, len(f.materials))
        _write_value(ws, row, 6, len(f.process_parameters))
        _write_value(ws, row, 7, f.source_image_index)
        _write_value(ws, row, 8, f.review_status)
        if f.review_status == "REVIEW_REQUIRED":
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = _CONFLICT_FILL
        row += 1


def _write_materials_sheet(wb: Workbook, entities: BusinessEntities) -> None:
    """配方明细：每条原料一行。"""
    ws = wb.create_sheet("配方明细")
    headers = ["公司", "产品/系列", "配方编号", "日期", "原料序号", "原料名称", "数量", "单位", "来源图片"]
    _write_header(ws, headers)

    row = 2
    for f in entities.formulas:
        company = _get_company_name(entities, f.company_id)
        product = _get_product_name(entities, f.product_id)
        no = f.formula_no_raw or f"未编号{f.formula_sequence}"
        for mi, m in enumerate(f.materials, 1):
            _write_value(ws, row, 1, company)
            _write_value(ws, row, 2, product)
            _write_value(ws, row, 3, no)
            _write_value(ws, row, 4, f.record_date.raw_value)
            _write_value(ws, row, 5, mi)
            _write_value(ws, row, 6, m.name.raw_value)
            _write_value(ws, row, 7, m.amount.raw_value)
            _write_value(ws, row, 8, m.unit.raw_value)
            _write_value(ws, row, 9, f.source_image_index)
            if m.amount.review_status == "CONFLICT":
                for col in range(1, 10):
                    ws.cell(row=row, column=col).fill = _CONFLICT_FILL
            row += 1


def _write_process_sheet(wb: Workbook, entities: BusinessEntities) -> None:
    """工艺参数：每条参数一行。"""
    ws = wb.create_sheet("工艺参数")
    headers = ["公司", "产品/系列", "配方编号", "日期", "参数序号", "参数名", "数值", "单位", "来源图片"]
    _write_header(ws, headers)

    row = 2
    for f in entities.formulas:
        company = _get_company_name(entities, f.company_id)
        product = _get_product_name(entities, f.product_id)
        no = f.formula_no_raw or f"未编号{f.formula_sequence}"
        for pi, p in enumerate(f.process_parameters, 1):
            _write_value(ws, row, 1, company)
            _write_value(ws, row, 2, product)
            _write_value(ws, row, 3, no)
            _write_value(ws, row, 4, f.record_date.raw_value)
            _write_value(ws, row, 5, pi)
            _write_value(ws, row, 6, p.name.raw_value)
            _write_value(ws, row, 7, p.value.raw_value)
            _write_value(ws, row, 8, p.unit.raw_value)
            _write_value(ws, row, 9, f.source_image_index)
            row += 1


def _write_review_sheet(wb: Workbook, entities: BusinessEntities) -> None:
    """识别审查：含公司原文/标准名、产品、配方ID、页面ID、bbox。"""
    ws = wb.create_sheet("识别审查")
    headers = [
        "公司原文", "公司标准名", "产品原文", "产品标准名",
        "配方ID", "配方编号", "页面ID", "来源图片", "record_bbox", "状态",
    ]
    _write_header(ws, headers)

    row = 2
    for f in entities.formulas:
        # 找页面公司原文
        page = next((p for p in entities.pages if p.page_id == f.page_id), None)
        company_raw = page.company.raw_value if page else ""
        company_std = _get_company_name(entities, f.company_id)
        product_std = _get_product_name(entities, f.product_id)

        # 找产品原文
        product_raw = ""
        if page:
            for section in page.product_sections:
                for fb in section.formula_blocks:
                    if fb.formula_id == f.formula_id:
                        product_raw = section.product_or_series.raw_value
                        break

        _write_value(ws, row, 1, company_raw)
        _write_value(ws, row, 2, company_std)
        _write_value(ws, row, 3, product_raw)
        _write_value(ws, row, 4, product_std)
        _write_value(ws, row, 5, f.formula_id)
        _write_value(
            ws,
            row,
            6,
            f.formula_no_raw or f"未编号{f.formula_sequence}",
        )
        _write_value(ws, row, 7, f.page_id)
        _write_value(ws, row, 8, f.source_image_index)
        _write_value(ws, row, 9, str(f.record_bbox) if f.record_bbox else "")
        _write_value(ws, row, 10, f.review_status)
        row += 1


def _write_correction_sheet_placeholder(wb: Workbook) -> None:
    """修正日志 Sheet（占位，实际数据由主导出器填充）。"""
    ws = wb.create_sheet("修正日志")
    headers = ["记录", "字段ID", "原值", "新值", "OCR", "VLM", "选择来源", "修改时间", "操作类型"]
    _write_header(ws, headers)
