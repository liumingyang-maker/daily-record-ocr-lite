"""Excel 导出系统：新建工作簿和固定模板映射。"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string

from .excel_safety import safe_excel_value
from .final_result import FinalResultError, FinalResultService
from .knowledge.path import resolve_knowledge_db_path
from .readiness import collect_unresolved_fields, evaluate_ready_gate
from .status import JobStatus
from .storage import JobStorage

logger = logging.getLogger(__name__)

MAX_AUTO_WIDTH = 50


class ExportError(RuntimeError):
    """导出错误。"""
    pass


class UnresolvedReviewError(ExportError):
    """Formal export was requested while reviewable fields remain."""


def export_job(job_id: str, storage: JobStorage | None = None) -> str:
    """
    导出任务结果为 Excel。

    返回输出文件名。
    """
    if storage is None:
        storage = JobStorage()

    job = storage.get_job(job_id)
    job_dir = storage.get_job_dir(job_id)
    demo_mode = bool(job.get("demo_mode"))
    try:
        result = FinalResultService(job_dir).load()
    except FinalResultError as exc:
        raise ExportError("正式导出要求 review/final_result.json。") from exc
    if not demo_mode and job.get("status") != JobStatus.READY:
        raise UnresolvedReviewError(
            f"正式导出仅允许 READY 任务；当前状态为 {job.get('status', '')}"
        )

    unresolved = collect_unresolved_fields(result)
    if not demo_mode:
        gate = evaluate_ready_gate(job, result, job_dir)
        if not gate.ready:
            raise UnresolvedReviewError("；".join(gate.reasons))
    if unresolved and not demo_mode:
        raise UnresolvedReviewError(
            f"仍有 {len(unresolved)} 个未确认字段，正式导出已阻止。"
        )

    prefix = "DEMO-" if demo_mode else ""
    output_name = f"{prefix}recognized-{job_id}.xlsx"
    output_relative = Path("export") / output_name
    output_path = job_dir / output_relative
    from .grouping.exporter import export_grouped_excel
    from .grouping.review import build_reviewed_business_entities

    entities = build_reviewed_business_entities(job_dir, job_id, result)
    job["status"] = JobStatus.EXPORTING
    job["status_message"] = "正在从正式 FinalResult 生成 Excel。"
    storage.save_job(job)
    export_grouped_excel(entities, output_path)
    logger.info("导出完成: %s", output_path)

    # 更新任务状态
    job["export_file"] = output_relative.as_posix()
    job["status"] = JobStatus.EXPORTED
    job["status_message"] = f"已导出: {output_name}"
    storage.save_job(job)

    return output_relative.as_posix()


def _iter_final_fields(result: dict[str, Any]):
    for page in result.get("pages", []):
        for section in page.get("product_sections", []):
            product = section.get("product_or_series")
            if isinstance(product, dict):
                yield product
            for formula in section.get("formulas", []):
                for name in ("record_date", "notes"):
                    field = formula.get(name)
                    if isinstance(field, dict):
                        yield field
                for material in formula.get("materials", []):
                    for name in ("name", "amount", "unit"):
                        field = material.get(name)
                        if isinstance(field, dict):
                            yield field
                for parameter in formula.get("process_parameters", []):
                    for name in ("name", "value", "unit"):
                        field = parameter.get(name)
                        if isinstance(field, dict):
                            yield field


def _get_sheet(wb: Workbook, sheet_name: str) -> Any:
    """获取或创建工作表。"""
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    return wb.create_sheet(title=sheet_name)


def _resolve_value(
    expr: str,
    item: dict[str, Any],
    root: dict[str, Any],
    parent: dict[str, Any] | None,
    index: int,
    parent_index: int,
) -> Any:
    """解析表达式值。"""
    if expr.startswith("literal:"):
        return expr[len("literal:"):]

    if expr == "$index":
        return index

    if expr == "$parent_index":
        return parent_index

    if expr.startswith("$root."):
        field = expr[len("$root."):]
        return _get_field(root, field)

    if expr.startswith("$parent."):
        if parent is None:
            return ""
        field = expr[len("$parent."):]
        return _get_field(parent, field)

    # 普通字段
    return _get_field(item, expr)


def _get_field(obj: dict[str, Any], field: str) -> Any:
    """获取字段值，列表/对象转为 JSON 字符串。"""
    value = obj.get(field, "")
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _write_cells(
    wb: Workbook, cells: list[dict[str, Any]], result: dict[str, Any]
) -> None:
    """写入单元格映射。"""
    for cell_cfg in cells:
        sheet_name = cell_cfg["sheet"]
        cell_ref = cell_cfg["cell"]
        expr = cell_cfg["value"]
        ws = _get_sheet(wb, sheet_name)
        value = _resolve_value(expr, {}, result, None, 0, 0)
        ws[cell_ref] = safe_excel_value(value)


def _write_table(
    wb: Workbook,
    table_cfg: dict[str, Any],
    result: dict[str, Any],
    is_template: bool,
) -> None:
    """写入表格映射。"""
    sheet_name = table_cfg["sheet"]
    source_key = table_cfg.get("source", "records")
    expand_key = table_cfg.get("expand", "")
    start_row = table_cfg.get("start_row", 1)
    include_header = table_cfg.get("include_header", True)
    auto_width = table_cfg.get("auto_width", False)
    columns = table_cfg.get("columns", [])

    ws = _get_sheet(wb, sheet_name)
    records = result.get(source_key, [])

    current_row = start_row

    # 写表头
    if include_header:
        for col_cfg in columns:
            col_letter = col_cfg["column"]
            header = col_cfg.get("header", "")
            col_idx = column_index_from_string(col_letter)
            cell = ws.cell(row=current_row, column=col_idx, value=header)
            if not is_template:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
                )
        current_row += 1

    # 写数据
    if expand_key:
        # 展开子数组
        for parent_idx, record in enumerate(records, 1):
            sub_items = record.get(expand_key, [])
            for item_idx, item in enumerate(sub_items, 1):
                for col_cfg in columns:
                    col_letter = col_cfg["column"]
                    expr = col_cfg["value"]
                    col_idx = column_index_from_string(col_letter)
                    value = _resolve_value(
                        expr, item, result, record, item_idx, parent_idx
                    )
                    ws.cell(
                        row=current_row,
                        column=col_idx,
                        value=safe_excel_value(value),
                    )
                current_row += 1
    else:
        # 直接写记录
        for item_idx, record in enumerate(records, 1):
            for col_cfg in columns:
                col_letter = col_cfg["column"]
                expr = col_cfg["value"]
                col_idx = column_index_from_string(col_letter)
                value = _resolve_value(
                    expr, record, result, None, item_idx, 0
                )
                ws.cell(
                    row=current_row,
                    column=col_idx,
                    value=safe_excel_value(value),
                )
            current_row += 1

    # 自动列宽
    if auto_width:
        _auto_column_width(ws, columns)


def _auto_column_width(ws: Any, columns: list[dict[str, Any]]) -> None:
    """自动调整列宽，限制最大宽度。"""
    for col_cfg in columns:
        col_letter = col_cfg["column"]
        max_len = 0
        for row in ws.iter_rows(
            min_col=column_index_from_string(col_letter),
            max_col=column_index_from_string(col_letter),
        ):
            for cell in row:
                if cell.value is not None:
                    cell_len = len(str(cell.value))
                    max_len = max(max_len, cell_len)
        # 中文字符占约2个宽度
        width = min(max_len + 2, MAX_AUTO_WIDTH)
        ws.column_dimensions[col_letter].width = max(width, 8)


def _write_audit_sheet(wb: Workbook, job_dir: Path, is_template: bool) -> None:
    """写入识别审查 Sheet。"""
    ws = _get_sheet(wb, "识别审查")
    headers = ["记录", "字段ID", "字段类型", "OCR", "OCR置信度", "VLM", "历史候选", "最终值", "来源", "状态"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        if not is_template:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # 读取融合结果
    fusion_path = job_dir / "fusion" / "result.json"
    if not fusion_path.exists():
        return

    try:
        fusion_data = json.loads(fusion_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return

    fields = fusion_data.get("fields", [])
    row = 2
    for f in fields:
        candidates = f.get("candidates", [])
        ocr_val = ""
        ocr_conf = ""
        vlm_val = ""
        history_val = ""
        for c in candidates:
            src = c.get("source", "")
            if src.startswith("ocr"):
                ocr_val = c.get("value", "")
                ocr_conf = str(c.get("confidence", ""))
            elif src == "vlm":
                vlm_val = c.get("value", "")
            elif src.startswith("history"):
                history_val = c.get("value", "")

        ws.cell(row=row, column=1, value=f.get("field_id", "").split("_")[0] if "_" in f.get("field_id", "") else "")
        ws.cell(row=row, column=2, value=f.get("field_id", ""))
        ws.cell(row=row, column=3, value=f.get("field_type", ""))
        ws.cell(row=row, column=4, value=ocr_val)
        ws.cell(row=row, column=5, value=ocr_conf)
        ws.cell(row=row, column=6, value=vlm_val)
        ws.cell(row=row, column=7, value=history_val)
        ws.cell(row=row, column=8, value=f.get("final_value", ""))
        ws.cell(row=row, column=9, value=f.get("final_source", ""))
        ws.cell(row=row, column=10, value=f.get("status", ""))

        # 冲突字段标红
        if f.get("status") == "CONFLICT":
            for col in range(1, 11):
                ws.cell(row=row, column=col).fill = PatternFill(
                    start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
                )
        row += 1


def _write_correction_sheet(wb: Workbook, job_id: str, is_template: bool) -> None:
    """写入修正日志 Sheet。"""
    ws = _get_sheet(wb, "修正日志")
    headers = ["记录", "字段ID", "原值", "新值", "OCR", "VLM", "选择来源", "修改时间"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        if not is_template:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # 读取修正日志
    try:
        from .knowledge.database import KnowledgeDB
        db = KnowledgeDB(resolve_knowledge_db_path())
        db.initialize()
        corrections = db.get_corrections_for_job(job_id)
        row = 2
        for c in corrections:
            ws.cell(row=row, column=1, value=c.get("record_id", ""))
            ws.cell(row=row, column=2, value=c.get("field_id", ""))
            ws.cell(row=row, column=3, value=c.get("old_value", ""))
            ws.cell(row=row, column=4, value=c.get("new_value", ""))
            ws.cell(row=row, column=5, value=c.get("ocr_value", ""))
            ws.cell(row=row, column=6, value=c.get("vlm_value", ""))
            ws.cell(row=row, column=7, value=c.get("chosen_source", ""))
            ws.cell(row=row, column=8, value=c.get("created_at", ""))
            row += 1
    except Exception as e:
        logger.warning("修正日志读取失败: %s", e)
