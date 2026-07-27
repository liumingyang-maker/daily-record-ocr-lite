"""FastAPI 页面与 API 路由。"""

from __future__ import annotations

import json
import logging
import os
import platform
import threading
import time
from contextlib import asynccontextmanager
from functools import lru_cache
from importlib.metadata import PackageNotFoundError
from importlib.metadata import version as package_version
from pathlib import Path
from urllib.parse import quote, urlparse

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from . import __version__
from .config import DATA_ROOT, get_config, load_recognition_config
from .contracts import normalize_legacy_result, validate_record_result
from .exporter import ExportError, export_job
from .final_result import (
    FinalResultError,
    FinalResultService,
    project_final_result,
)
from .presentation import present_stored_job
from .readiness import (
    collect_unresolved_fields,
    evaluate_content_gate,
    evaluate_ready_gate,
    validate_final_result_contract,
)
from .review_editor import ReviewEditor, ReviewInputError, ReviewVersionConflict
from .review_state import ReviewStateStore, confirm_formula
from .review_view import build_review_view
from .settings import SettingsService
from .status import JobStatus, SetupState
from .storage import JobStorage, write_json_atomic
from .upload_options import VALID_ROTATIONS, normalize_rotations

logger = logging.getLogger(__name__)
_model_install_guard = threading.Lock()
_model_install_thread: threading.Thread | None = None


@asynccontextmanager
async def lifespan(app: FastAPI):
    """应用生命周期：启动时初始化 OCR 管理器和后台队列。"""
    try:
        from .ocr.manager import OCRModelManager
        rec_cfg = load_recognition_config()
        ocr_cfg = rec_cfg.get("ocr", {})
        mgr = OCRModelManager()
        mgr.configure({
            "enabled": str(ocr_cfg.get("enabled", "true")).lower() in ("true", "1", "yes"),
            "provider": ocr_cfg.get("provider", ""),
            "device": ocr_cfg.get("device", "cpu"),
            "tier": ocr_cfg.get("tier", "medium"),
            "minimum_score": float(ocr_cfg.get("minimum_score", 0.45)),
            "use_textline_orientation": ocr_cfg.get("use_textline_orientation", True),
        })
        logger.info("OCR 管理器已配置: provider=%s", ocr_cfg.get("provider", ""))
    except Exception as e:
        logger.warning("OCR 管理器初始化失败（降级运行）: %s", e)

    # 启动后台任务队列
    from .jobs import get_task_queue
    queue = get_task_queue()
    queue.set_handler(_process_job_background)
    await queue.start()

    yield

    await queue.stop()


async def _process_job_background(job_id: str) -> None:
    """后台处理单个任务：双引擎 pipeline。失败时标记 FAILED/DEGRADED，不回退旧 pipeline。"""
    storage = _get_storage()
    try:
        from .pipeline_v2 import analyze_job_v2
        await analyze_job_v2(job_id, storage)
    except Exception as e:
        logger.error("任务 %s 双引擎识别失败: %s", job_id, e)
        try:
            job = storage.get_job(job_id)
            if job.get("status") not in {
                JobStatus.FAILED,
                JobStatus.FAILED_SCHEMA,
            }:
                job["status"] = JobStatus.DEGRADED
                job["error"] = f"双引擎识别失败: {e}"
                job["status_message"] = (
                    "双引擎识别未完成，结果不可靠。请检查 OCR/VLM 配置后重新识别。"
                )
                storage.save_job(job)
        except Exception:
            logger.exception("任务 %s 状态更新失败", job_id)


app = FastAPI(title="daily-record-ocr-lite", version=__version__, lifespan=lifespan)

# 静态文件和模板
_BASE_DIR = Path(__file__).resolve().parent
app.mount("/static", StaticFiles(directory=str(_BASE_DIR / "static")), name="static")


def _base_template_context(request: Request) -> dict:
    return {
        "demo_mode": _demo_enabled(),
        "setup_state": _recognition_gate(),
        "version": __version__,
    }


templates = Jinja2Templates(
    directory=str(_BASE_DIR / "templates"),
    context_processors=[_base_template_context],
)


def _get_storage() -> JobStorage:
    return JobStorage()


def _get_settings() -> SettingsService:
    return SettingsService(DATA_ROOT)


def _demo_enabled() -> bool:
    if os.environ.get("DEMO_MODE", "").strip().lower() in {
        "1",
        "true",
        "yes",
        "on",
    }:
        return True
    return bool(_get_settings().load().get("demo_mode"))


def _recognition_gate() -> str:
    if not _desktop_models_ready():
        return "SETUP_REQUIRED"
    return str(_get_settings().status()["state"])


def _desktop_models_ready() -> bool:
    from .platform_paths import runtime_kind

    return _model_ready_for_root(str(DATA_ROOT), runtime_kind())


@lru_cache(maxsize=4)
def _model_ready_for_root(data_root: str, install_kind: str) -> bool:
    if install_kind != "desktop":
        return True
    from .model_packages import model_package_status

    return model_package_status(Path(data_root))["status"] == "READY"


def _setup_system_info() -> dict:
    from .ocr.manager import OCRModelManager

    def installed_version(distribution: str) -> str:
        try:
            return package_version(distribution)
        except PackageNotFoundError:
            return "未安装"

    ocr = OCRModelManager().get_status()
    return {
        "app_version": __version__,
        "python": platform.python_version(),
        "platform": f"{platform.system()} {platform.release()} ({platform.machine()})",
        "paddle": installed_version("paddlepaddle"),
        "paddleocr": installed_version("paddleocr"),
        "ocr": ocr,
    }


def _guard_local_json_write(request: Request) -> None:
    host = request.headers.get("host", "").split(":", 1)[0].lower()
    if host not in {"127.0.0.1", "localhost", "::1", "testserver"}:
        raise HTTPException(status_code=403, detail="配置写入只允许本机访问。")
    origin = request.headers.get("origin")
    if origin:
        origin_host = (urlparse(origin).hostname or "").lower()
        if origin_host not in {"127.0.0.1", "localhost", "::1", "testserver"}:
            raise HTTPException(status_code=403, detail="拒绝跨站配置写入。")
    if not request.headers.get("content-type", "").lower().startswith("application/json"):
        raise HTTPException(status_code=415, detail="配置写入必须使用 JSON。")


def _unresolved_final_fields(final: dict) -> list[str]:
    return collect_unresolved_fields(final)


def _apply_ready_gate(job: dict, final: dict, job_dir: Path) -> list[str]:
    gate = evaluate_ready_gate(job, final, job_dir)
    if gate.ready:
        job["status"] = JobStatus.READY
        job["status_message"] = "所有 READY Gate 均已通过。"
    else:
        job["status"] = JobStatus.REVIEW_REQUIRED
        job["status_message"] = "；".join(gate.reasons)
    return gate.unresolved_fields


# ─── 健康检查 ───────────────────────────────────────────────


@app.get("/health")
async def health():
    return {
        "status": "ok",
        "app": "daily-record-ocr-lite",
        "version": __version__,
        "setup_state": _recognition_gate(),
    }


@app.get("/api/engine-status")
async def engine_status():
    from .jobs import get_task_queue
    from .ocr.manager import OCRModelManager

    ocr_mgr = OCRModelManager()
    cfg = get_config()
    vision_cfg = cfg.vision
    return {
        "ocr": ocr_mgr.get_status(),
        "vision": {
            "configured": _recognition_gate() == SetupState.READY_FOR_RECOGNITION,
            "provider": vision_cfg.get("provider", ""),
            "model": vision_cfg.get("model", ""),
            "api_key_configured": bool(vision_cfg.get("api_key")),
        },
        "queue": get_task_queue().get_status(),
        "setup_state": _recognition_gate(),
        "demo_mode": _demo_enabled(),
        "version": __version__,
    }


# ─── 首页 ───────────────────────────────────────────────────


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    if _recognition_gate() == SetupState.SETUP_REQUIRED:
        return RedirectResponse(url="/setup", status_code=302)
    storage = _get_storage()
    jobs = [
        present_stored_job(job, storage.get_job_dir(job["id"]))
        for job in storage.list_jobs()[:5]
    ]
    return templates.TemplateResponse(
        request,
        "index.html",
        {
            "jobs": jobs,
            "setup_state": _recognition_gate(),
            "demo_mode": _demo_enabled(),
        },
    )


@app.get("/jobs", response_class=HTMLResponse)
async def job_records(request: Request):
    """List recognition records using business language."""
    storage = _get_storage()
    jobs = [
        present_stored_job(job, storage.get_job_dir(job["id"]))
        for job in storage.list_jobs()
    ]
    return templates.TemplateResponse(request, "jobs.html", {"jobs": jobs})


@app.get("/setup", response_class=HTMLResponse)
async def setup_page(request: Request):
    return templates.TemplateResponse(
        request,
        "setup.html",
        {
            "settings": _get_settings().public_settings(),
            "status": _get_settings().status(),
            "system": _setup_system_info(),
            "version": __version__,
        },
    )


@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request):
    return templates.TemplateResponse(
        request,
        "settings.html",
        {
            "settings": _get_settings().public_settings(),
            "status": _get_settings().status(),
            "version": __version__,
        },
    )


@app.get("/api/settings/status")
async def settings_status():
    return _get_settings().status()


@app.get("/api/settings/public")
async def settings_public():
    return _get_settings().public_settings()


@app.get("/api/updates/stable")
def stable_update_status():
    from .updates import check_for_stable_update

    return check_for_stable_update(__version__)


@app.post("/api/desktop/import-data")
async def import_existing_desktop_data(request: Request):
    from .desktop_migration import (
        InvalidMigrationSource,
        MigrationConfirmationRequired,
        import_legacy_data,
    )

    _guard_local_json_write(request)
    body = await request.json()
    source_path = str(body.get("source_path", "")).strip()
    if not source_path:
        raise HTTPException(status_code=400, detail="请选择现有项目或 data 目录")
    try:
        report = import_legacy_data(
            Path(source_path),
            DATA_ROOT,
            confirm_non_empty=bool(body.get("confirm_non_empty", False)),
        )
    except MigrationConfirmationRequired as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except InvalidMigrationSource as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return report.to_safe_dict()


@app.get("/api/setup/models/status")
def desktop_model_status():
    from .model_packages import model_package_status

    state_path = DATA_ROOT / "model_install_state.json"
    if _model_install_thread is not None and _model_install_thread.is_alive():
        try:
            state = json.loads(state_path.read_text(encoding="utf-8"))
            if isinstance(state, dict):
                return state
        except (OSError, json.JSONDecodeError):
            pass
    return model_package_status(DATA_ROOT)


@app.post("/api/setup/models/install", status_code=202)
async def install_desktop_models(request: Request):
    global _model_install_thread

    _guard_local_json_write(request)
    state_path = DATA_ROOT / "model_install_state.json"

    def save_state(payload: dict[str, object]) -> None:
        allowed = {
            key: payload[key]
            for key in (
                "status",
                "state",
                "model",
                "index",
                "package_count",
                "downloaded_bytes",
                "total_bytes",
                "ready_count",
                "error_category",
            )
            if key in payload
        }
        write_json_atomic(state_path, allowed)

    def worker() -> None:
        from .model_packages import install_model_packages

        try:
            result = install_model_packages(
                DATA_ROOT,
                progress=lambda event: save_state({"status": "INSTALLING", **event}),
            )
            _model_ready_for_root.cache_clear()
            save_state(result)
        except Exception as exc:
            logger.error(
                "Desktop OCR model installation failed: %s",
                type(exc).__name__,
            )
            save_state({"status": "FAILED", "error_category": "MODEL_INSTALL_FAILED"})

    with _model_install_guard:
        if _model_install_thread is None or not _model_install_thread.is_alive():
            save_state({"status": "INSTALLING", "state": "STARTING"})
            _model_install_thread = threading.Thread(
                target=worker,
                name="desktop-model-installer",
                daemon=True,
            )
            _model_install_thread.start()
    return {"status": "INSTALLING"}


@app.put("/api/settings/vision")
async def configure_vision(request: Request):
    _guard_local_json_write(request)
    result = _get_settings().configure_vision(await request.json())
    return {"status": _get_settings().status(), "settings": result}


@app.put("/api/settings/ocr")
async def configure_ocr(request: Request):
    _guard_local_json_write(request)
    result = _get_settings().configure_ocr(await request.json())
    from .ocr.manager import OCRModelManager

    OCRModelManager().configure(_get_settings().effective_settings()["ocr"])
    return {"status": _get_settings().status(), "settings": result}


@app.post("/api/settings/demo/enable")
async def enable_demo(request: Request):
    _guard_local_json_write(request)
    return _get_settings().enable_demo()


@app.post("/api/settings/demo/disable")
async def disable_demo(request: Request):
    _guard_local_json_write(request)
    return _get_settings().disable_demo()


@app.post("/api/settings/test-vision")
async def test_vision_settings(request: Request):
    _guard_local_json_write(request)
    vision = _get_settings().effective_settings()["vision"]
    if _demo_enabled():
        return {
            "status": "DEMO_MODE",
            "provider": vision.get("provider", ""),
            "model": vision.get("model", ""),
            "latency_ms": 0,
            "http_status": None,
            "message": "演示模式使用固定结果；未发起真实视觉模型请求。",
            "vision_capability": False,
            "json_response_capability": False,
            "strict_json_capability": False,
            "response_preview": "",
            "error_category": None,
        }
    from .pipeline_v2 import build_vision_provider
    from .vision.probe import (
        PROBE_SCHEMA,
        PROBE_SYSTEM_PROMPT,
        PROBE_USER_PROMPT,
        create_probe_image,
        evaluate_probe_response,
        safe_response_preview,
    )

    provider = build_vision_provider(vision)
    _get_settings().data_dir.mkdir(parents=True, exist_ok=True)
    test_image = _get_settings().data_dir / "connection-test.png"
    create_probe_image(test_image)
    started = time.monotonic()
    try:
        raw = await provider.analyze(
            [test_image],
            PROBE_SYSTEM_PROMPT,
            PROBE_USER_PROMPT,
            PROBE_SCHEMA,
        )
    except Exception as exc:
        _get_settings().record_health("vision", False)
        message = str(exc).lower()
        if "401" in message or "403" in message:
            category = "AUTHENTICATION"
        elif "timeout" in message or "超时" in message:
            category = "TIMEOUT"
        elif "json" in message or "schema" in message:
            category = "JSON_CAPABILITY"
        else:
            category = "PROVIDER_OR_NETWORK"
        return JSONResponse(
            status_code=502,
            content={
                "status": "ERROR",
                "provider": vision.get("provider", ""),
                "model": vision.get("model", ""),
                "latency_ms": int((time.monotonic() - started) * 1000),
                "http_status": None,
                "vision_capability": False,
                "json_response_capability": False,
                "strict_json_capability": False,
                "response_preview": "",
                "error_category": category,
                "error_type": type(exc).__name__,
            },
        )
    finally:
        test_image.unlink(missing_ok=True)
    capabilities = evaluate_probe_response(raw)
    _get_settings().record_health("vision", capabilities.vision_capability)
    if capabilities.vision_capability:
        status = "OK"
        error_category = None
    elif capabilities.json_response_capability:
        status = "FAILED_VISION_CAPABILITY"
        error_category = "VISION_CAPABILITY"
    else:
        status = "FAILED_JSON_CAPABILITY"
        error_category = "JSON_CAPABILITY"
    return {
        "status": status,
        "provider": vision.get("provider", ""),
        "model": vision.get("model", ""),
        "latency_ms": int((time.monotonic() - started) * 1000),
        "http_status": 200,
        "vision_capability": capabilities.vision_capability,
        "json_response_capability": capabilities.json_response_capability,
        "strict_json_capability": capabilities.strict_json_capability,
        "response_preview": safe_response_preview(raw, vision.get("api_key", "")),
        "error_category": error_category,
    }


@app.post("/api/settings/test-ocr")
async def test_ocr_settings(request: Request):
    _guard_local_json_write(request)
    if not _desktop_models_ready():
        return JSONResponse(
            status_code=409,
            content={
                "status": "SETUP_REQUIRED",
                "error_category": "MODEL_DOWNLOAD_REQUIRED",
                "message": "请先下载并校验 OCR 模型",
            },
        )
    from PIL import Image, ImageDraw, ImageFont

    from .ocr.manager import OCRModelManager
    from .ocr.overlay import generate_overlay

    manager = OCRModelManager()
    manager.configure(_get_settings().effective_settings()["ocr"])
    setup_dir = _get_settings().data_dir / "setup"
    setup_dir.mkdir(parents=True, exist_ok=True)
    test_image = setup_dir / "ocr-test.png"
    overlay = setup_dir / "ocr-overlay.jpg"
    image = Image.new("RGB", (720, 260), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 42)
    except OSError:
        font = ImageFont.load_default()
    draw.text((40, 35), "PA66 60kg", fill="black", font=font)
    draw.text((40, 105), "GF30 30kg", fill="black", font=font)
    draw.text((40, 175), "Process 50Hz", fill="black", font=font)
    image.save(test_image)
    try:
        page = await manager.recognize_async(test_image)
        generate_overlay(test_image, page, overlay)
    except Exception:
        _get_settings().record_health("ocr", False)
        raise
    finally:
        test_image.unlink(missing_ok=True)
    payload = {
        "status": "OK" if page.tokens else "FAILED_NO_TOKENS",
        "provider": page.provider,
        "model": page.model,
        "models": manager.get_status(),
        "token_count": len(page.tokens),
        "token_preview": [
            {"text": token.text, "confidence": token.confidence}
            for token in page.tokens[:8]
        ],
        "elapsed_ms": page.elapsed_ms,
        "overlay_url": f"/api/settings/ocr-overlay?v={int(time.time())}",
    }
    if not page.tokens:
        _get_settings().record_health("ocr", False)
        return JSONResponse(status_code=422, content=payload)
    _get_settings().record_health("ocr", True)
    return payload


@app.get("/api/settings/ocr-overlay")
async def get_ocr_test_overlay():
    overlay = _get_settings().data_dir / "setup" / "ocr-overlay.jpg"
    if not overlay.exists():
        raise HTTPException(status_code=404, detail="尚未运行 OCR 测试。")
    return FileResponse(overlay, media_type="image/jpeg")


# ─── 创建任务并识别 ─────────────────────────────────────────


@app.post("/jobs")
async def create_job(
    files: list[UploadFile] = File(...),
    rotation: str = Form("auto"),
    rotation_manifest: str = Form(""),
):
    import io as _io
    import shutil

    from PIL import Image as PILImage

    cfg = get_config()
    storage = _get_storage()
    if _recognition_gate() == SetupState.SETUP_REQUIRED:
        raise HTTPException(
            status_code=409,
            detail="SETUP_REQUIRED：请先配置视觉模型，或显式进入演示模式。",
        )

    # 校验旋转值
    if rotation not in VALID_ROTATIONS:
        raise HTTPException(
            status_code=400,
            detail=f"无效的旋转设置: {rotation}。允许: {', '.join(sorted(VALID_ROTATIONS))}",
        )

    # 校验文件
    if not files:
        raise HTTPException(status_code=400, detail="请至少上传一张图片。")

    max_bytes = cfg.max_upload_mb * 1024 * 1024
    total_size = 0

    # 先读取并验证所有文件（在创建任务之前）
    validated_files: list[tuple[str, bytes]] = []
    for f in files:
        if not f.filename:
            raise HTTPException(status_code=400, detail="文件名不能为空。")
        ext = Path(f.filename).suffix.lower()
        if ext not in cfg.allowed_extensions:
            raise HTTPException(
                status_code=400,
                detail=f"不支持的文件类型: {ext}。允许: {', '.join(cfg.allowed_extensions)}",
            )
        content = await f.read()
        if not content:
            raise HTTPException(status_code=400, detail=f"文件为空: {f.filename}")
        total_size += len(content)
        if total_size > max_bytes:
            raise HTTPException(
                status_code=400,
                detail=f"上传文件总大小超过限制 ({cfg.max_upload_mb}MB)。",
            )
        # Pillow 预验证：确认是有效图片
        try:
            img = PILImage.open(_io.BytesIO(content))
            img.verify()
        except Exception:
            raise HTTPException(
                status_code=400,
                detail=f"文件不是可识别的图片: {f.filename}",
            )
        validated_files.append((f.filename, content))

    try:
        rotations = normalize_rotations(
            rotation_manifest,
            count=len(validated_files),
            fallback=rotation,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    # 创建任务
    job = storage.create_job(rotation=rotation, rotations=rotations)
    job_id = job["id"]
    job["demo_mode"] = _demo_enabled()

    # 保存上传文件
    try:
        for i, (filename, content) in enumerate(validated_files, 1):
            img_info = storage.save_upload(job_id, i, filename, content)
            job["images"].append(img_info)
        storage.save_job(job)
    except Exception as e:
        # 清理失败的任务目录
        try:
            shutil.rmtree(storage.get_job_dir(job_id), ignore_errors=True)
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=f"文件保存失败: {e}")

    # 提交到后台队列（非阻塞）
    from .jobs import get_task_queue
    queue = get_task_queue()
    await queue.submit(job_id)

    return RedirectResponse(url=f"/jobs/{job_id}", status_code=303)


# ─── 任务详情 ───────────────────────────────────────────────


@app.get("/jobs/{job_id}", response_class=HTMLResponse)
async def job_detail(job_id: str, request: Request):
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    presented = present_stored_job(job, storage.get_job_dir(job_id))
    job = {
        **job,
        "title": f"{presented['customer']} · {presented['product']}",
        "user_status": presented["status"]["label"],
        "progress_step": presented["progress_step"],
        "progress_label": presented["progress_label"],
    }
    return templates.TemplateResponse(
        request,
        "job.html",
        {
            "job": job,
        },
    )


@app.get("/jobs/{job_id}/result")
async def job_result_page(job_id: str):
    """Keep old bookmarks working while using one review workspace."""
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")
    return RedirectResponse(url=f"/jobs/{job_id}#review", status_code=303)


# ─── 重新识别 ───────────────────────────────────────────────


@app.post("/jobs/{job_id}/analyze")
async def reanalyze(job_id: str):
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    # 提交到后台队列重新识别
    from .jobs import get_task_queue
    queue = get_task_queue()
    await queue.submit(job_id)

    return RedirectResponse(url=f"/jobs/{job_id}", status_code=303)


@app.post("/api/jobs/{job_id}/mode")
async def switch_mode(job_id: str, request: Request):
    """切换识别模式（freeform/template/auto）并重新识别。"""
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    mode = body.get("mode", "auto")
    valid_modes = {"freeform", "template", "auto"}
    if mode not in valid_modes:
        raise HTTPException(
            status_code=400,
            detail=f"无效模式: {mode}。允许: {', '.join(sorted(valid_modes))}",
        )

    job["recognition_mode"] = mode
    job["status"] = JobStatus.UPLOADED
    job["status_message"] = f"已切换为 {mode} 模式，等待重新识别。"
    storage.save_job(job)

    # 提交到后台队列
    from .jobs import get_task_queue
    queue = get_task_queue()
    await queue.submit(job_id)

    return {"status": JobStatus.UPLOADED, "mode": mode}


@app.post("/api/jobs/{job_id}/reanalyze")
async def reanalyze_api(job_id: str):
    """重新识别（API 版本，提交到后台队列）。"""
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    job["status"] = JobStatus.UPLOADED
    job["status_message"] = "等待重新识别..."
    storage.save_job(job)

    from .jobs import get_task_queue
    queue = get_task_queue()
    await queue.submit(job_id)

    return {"status": JobStatus.UPLOADED, "message": "已提交重新识别"}


# ─── 保存人工修改 ───────────────────────────────────────────


@app.post("/api/jobs/{job_id}/result")
async def save_result(job_id: str, request: Request):
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.body()
    try:
        result = json.loads(body)
    except json.JSONDecodeError as e:
        return JSONResponse(
            status_code=400,
            content={"status": "ERROR", "validation_errors": [f"JSON 格式错误: {e}"]},
        )

    if not isinstance(result, dict):
        return JSONResponse(
            status_code=400,
            content={"status": "ERROR", "validation_errors": ["结果必须是 JSON 对象。"]},
        )

    normalized = normalize_legacy_result(result, job_id)
    report = validate_record_result(normalized)
    if report.issues:
        job["validation_errors"] = [
            issue.to_dict() for issue in report.issues
        ]
        job["status"] = (
            JobStatus.FAILED_SCHEMA
            if report.fatal
            else JobStatus.REVIEW_REQUIRED
        )
        job["status_message"] = "结果不符合 record-v1，未覆盖正式 FinalResult。"
        storage.save_job(job)
        return JSONResponse(
            status_code=422,
            content={
                "status": job["status"],
                "validation_errors": job["validation_errors"],
            },
        )

    job_dir = storage.get_job_dir(job_id)
    final_contract = validate_final_result_contract(normalized)
    if normalized.get("job_id") == job_id and not final_contract.issues:
        final = normalized
    else:
        final = project_final_result(job_id, normalized, {"fields": []})
    final["recognition_run_id"] = str(job.get("recognition_run_id", ""))
    FinalResultService(job_dir).replace(final)
    job["validation_errors"] = []
    job["final_result_run_id"] = job.get("recognition_run_id")
    unresolved = _apply_ready_gate(job, final, job_dir)
    storage.save_job(job)

    return {
        "status": job["status"],
        "validation_errors": [],
        "unresolved_fields": unresolved,
    }


# ─── 导出 ───────────────────────────────────────────────────


@app.post("/jobs/{job_id}/export")
async def export(job_id: str):
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    try:
        export_job(job_id, storage)
    except ExportError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出失败: {e}")

    return RedirectResponse(url=f"/jobs/{job_id}", status_code=303)


@app.post("/api/jobs/{job_id}/export")
async def export_api(job_id: str):
    """导出 Excel（API 版本，返回 JSON）。"""
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    try:
        filename = export_job(job_id, storage)
    except ExportError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出失败: {e}")

    return {
        "status": JobStatus.EXPORTED,
        "filename": filename,
        "download_url": f"/jobs/{job_id}/files/{filename}",
    }


# ─── 文件查看和下载 ─────────────────────────────────────────


@app.get("/jobs/{job_id}/files/{filename:path}")
async def download_file(job_id: str, filename: str):
    storage = _get_storage()
    try:
        file_path = storage.get_file_path(job_id, filename)
    except (FileNotFoundError, ValueError) as e:
        raise HTTPException(status_code=404, detail=str(e))

    # 判断是图片还是其他文件
    ext = file_path.suffix.lower()
    mime_map = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
    }
    if ext in mime_map:
        return FileResponse(file_path, media_type=mime_map[ext])
    elif ext == ".xlsm":
        return FileResponse(
            file_path,
            media_type="application/vnd.ms-excel.sheet.macroEnabled.12",
            filename=Path(filename).name,
        )
    elif ext == ".xlsx":
        return FileResponse(
            file_path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=Path(filename).name,
        )
    else:
        return FileResponse(file_path, filename=Path(filename).name)


# ─── 业务配方审查 API ────────────────────────────────────────


def _review_resources(job_id: str) -> tuple[JobStorage, dict, ReviewEditor, ReviewStateStore]:
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")
    job_dir = storage.get_job_dir(job_id)
    return storage, job, ReviewEditor(job_dir), ReviewStateStore(job_dir)


def _review_error(exc: Exception) -> HTTPException:
    if isinstance(exc, ReviewVersionConflict):
        return HTTPException(status_code=409, detail=str(exc))
    status_code = 404 if "不存在" in str(exc) else 422
    return HTTPException(status_code=status_code, detail=str(exc))


def _version_from(body: dict) -> str:
    version = str(body.get("version", ""))
    if not version:
        raise HTTPException(status_code=422, detail="缺少页面版本，请刷新后重试。")
    return version


def _saved(editor: ReviewEditor, message: str, **extra) -> dict:
    return {
        "version": str(editor.load()["updated_at"]),
        "saved": True,
        "message": message,
        **extra,
    }


@app.get("/api/jobs/{job_id}/review")
async def get_review(job_id: str):
    _storage, job, editor, state = _review_resources(job_id)
    try:
        final = editor.load()
    except FinalResultError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    view = build_review_view(job, final, state.confirmation_map(final))
    view["version"] = str(final["updated_at"])
    return view


@app.patch("/api/jobs/{job_id}/review/groups/{group_id}")
async def update_review_identity(job_id: str, group_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    try:
        editor.update_identity(
            group_id,
            customer=str(body.get("customer", "")),
            product=str(body.get("product", "")),
            expected_version=_version_from(body),
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "客户和产品已保存")


@app.post("/api/jobs/{job_id}/review/groups/{group_id}/formulas")
async def add_review_formula(job_id: str, group_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    try:
        formula_id = editor.add_formula(group_id, expected_version=_version_from(body))
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "已增加配方", formula_id=formula_id)


@app.patch("/api/jobs/{job_id}/review/formulas/{formula_id}")
async def update_review_formula(job_id: str, formula_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    changes = {key: value for key, value in body.items() if key != "version"}
    try:
        editor.update_formula(
            formula_id,
            changes,
            expected_version=_version_from(body),
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "配方信息已保存")


@app.delete("/api/jobs/{job_id}/review/formulas/{formula_id}")
async def delete_review_formula(job_id: str, formula_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    if body.get("confirmed") is not True:
        raise HTTPException(status_code=422, detail="请先确认删除整条配方。")
    try:
        editor.delete_formula(formula_id, expected_version=_version_from(body))
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "配方已删除，可立即撤销", undo_available=True)


@app.post("/api/jobs/{job_id}/review/formulas/{formula_id}/materials")
async def add_review_material(job_id: str, formula_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    values = {key: body.get(key, "") for key in ("name", "amount", "unit")}
    try:
        material_id = editor.add_material(
            formula_id, values, expected_version=_version_from(body)
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "材料已增加", material_id=material_id)


@app.patch("/api/jobs/{job_id}/review/formulas/{formula_id}/materials/{material_id}")
async def update_review_material(
    job_id: str, formula_id: str, material_id: str, request: Request
):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    changes = {key: value for key, value in body.items() if key != "version"}
    try:
        editor.update_material(
            formula_id,
            material_id,
            changes,
            expected_version=_version_from(body),
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "材料已保存")


@app.post("/api/jobs/{job_id}/review/formulas/{formula_id}/materials/reorder")
async def reorder_review_materials(job_id: str, formula_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    material_ids = body.get("material_ids", [])
    if not isinstance(material_ids, list):
        raise HTTPException(status_code=422, detail="材料顺序必须是数组。")
    try:
        editor.reorder_materials(
            formula_id,
            [str(item) for item in material_ids],
            expected_version=_version_from(body),
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "材料顺序已保存")


@app.delete("/api/jobs/{job_id}/review/formulas/{formula_id}/materials/{material_id}")
async def delete_review_material(
    job_id: str, formula_id: str, material_id: str, request: Request
):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    try:
        editor.delete_material(
            formula_id, material_id, expected_version=_version_from(body)
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "材料已删除，可立即撤销", undo_available=True)


@app.post("/api/jobs/{job_id}/review/formulas/{formula_id}/process")
async def add_review_process(job_id: str, formula_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    values = {key: body.get(key, "") for key in ("name", "value", "unit")}
    try:
        parameter_id = editor.add_process_parameter(
            formula_id, values, expected_version=_version_from(body)
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "工艺已增加", parameter_id=parameter_id)


@app.patch("/api/jobs/{job_id}/review/formulas/{formula_id}/process/{parameter_id}")
async def update_review_process(
    job_id: str, formula_id: str, parameter_id: str, request: Request
):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    changes = {key: value for key, value in body.items() if key != "version"}
    try:
        editor.update_process_parameter(
            formula_id,
            parameter_id,
            changes,
            expected_version=_version_from(body),
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "工艺已保存")


@app.delete("/api/jobs/{job_id}/review/formulas/{formula_id}/process/{parameter_id}")
async def delete_review_process(
    job_id: str, formula_id: str, parameter_id: str, request: Request
):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    try:
        editor.delete_process_parameter(
            formula_id, parameter_id, expected_version=_version_from(body)
        )
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "工艺已删除，可立即撤销", undo_available=True)


@app.post("/api/jobs/{job_id}/review/undo")
async def undo_review_delete(job_id: str, request: Request):
    _storage, _job, editor, _state = _review_resources(job_id)
    body = await request.json()
    try:
        editor.undo_last_delete(expected_version=_version_from(body))
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return _saved(editor, "已恢复刚才删除的内容")


@app.post("/api/jobs/{job_id}/review/formulas/{formula_id}/confirm")
async def confirm_review_formula(job_id: str, formula_id: str, request: Request):
    _storage, _job, editor, state = _review_resources(job_id)
    body = await request.json()
    try:
        result = confirm_formula(editor, state, formula_id, _version_from(body))
    except (ReviewInputError, ReviewVersionConflict) as exc:
        raise _review_error(exc) from exc
    return {
        **_saved(editor, "这条配方已确认"),
        "confirmed": result["confirmed"],
    }


# ─── 字段级兼容 API ─────────────────────────────────────────


@app.get("/api/jobs/{job_id}")
async def get_job_api(job_id: str):
    """获取任务状态（JSON API，供前端轮询）。"""
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")
    return job


async def _perform_recheck(
    job_id: str,
    field_ids: list[str],
    storage: JobStorage,
) -> dict:
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    from .ocr.manager import OCRModelManager
    from .pipeline_v2 import build_vision_provider
    from .review.recheck import RecheckError, recheck_fields

    vision_provider = None
    if not _demo_enabled():
        try:
            vision_provider = build_vision_provider(get_config().vision)
        except Exception as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
    try:
        result = await recheck_fields(
            job_id,
            storage.get_job_dir(job_id),
            job,
            field_ids,
            OCRModelManager(),
            vision_provider,
        )
    except (FinalResultError, RecheckError) as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc

    final = FinalResultService(storage.get_job_dir(job_id)).load()
    unresolved = _apply_ready_gate(job, final, storage.get_job_dir(job_id))
    storage.save_job(job)
    result["job_status"] = job["status"]
    result["unresolved_fields"] = unresolved
    return result


@app.post("/api/jobs/{job_id}/fields/recheck")
async def recheck_fields_batch(job_id: str, request: Request):
    body = await request.json()
    field_ids = body.get("field_ids", [])
    if not isinstance(field_ids, list):
        raise HTTPException(status_code=422, detail="field_ids 必须是数组。")
    return await _perform_recheck(job_id, field_ids, _get_storage())


@app.post("/api/jobs/{job_id}/fields/{field_id}")
async def update_field(job_id: str, field_id: str, request: Request):
    """人工修改单个字段值。"""
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    new_value = body.get("value", "")
    chosen_source = body.get("source", "manual")

    job_dir = storage.get_job_dir(job_id)
    service = FinalResultService(job_dir)
    try:
        if new_value == "" and chosen_source != "manual":
            updated = service.choose_candidate(field_id, str(chosen_source))
        else:
            updated = service.update_field(
                field_id,
                str(new_value),
                str(chosen_source),
            )
    except FinalResultError as exc:
        message = str(exc)
        status_code = 404 if "字段不存在" in message else 409
        raise HTTPException(status_code=status_code, detail=message) from exc

    # 记录修正日志
    try:
        from .knowledge.database import KnowledgeDB
        db = KnowledgeDB(DATA_ROOT / "knowledge.sqlite3")
        db.initialize()
        from .review.corrections import CorrectionService
        svc = CorrectionService(db)
        svc.record_correction(
            job_id=job_id,
            field_id=field_id,
            field_type=body.get("field_type", ""),
            old_value="",
            new_value=new_value,
            chosen_source=chosen_source,
        )
    except Exception as e:
        logger.warning("修正日志记录失败: %s", e)

    final = service.load()
    _apply_ready_gate(job, final, job_dir)
    storage.save_job(job)
    return {
        "status": updated["status"],
        "field_id": field_id,
        "value": updated["value"],
        "job_status": job["status"],
    }


def _finalize_job(job_id: str) -> dict:
    """Append confirmed formulas, bind a receipt, then cross the READY gate."""
    storage = _get_storage()
    try:
        job = storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    job_dir = storage.get_job_dir(job_id)
    try:
        final = FinalResultService(job_dir).load()
    except FinalResultError as exc:
        job["status"] = JobStatus.REVIEW_REQUIRED
        job["status_message"] = f"无法确认：{exc}"
        storage.save_job(job)
        raise HTTPException(status_code=409, detail=job["status_message"]) from exc

    content = evaluate_content_gate(job, final, job_dir)
    if not content.ready:
        job["status"] = JobStatus.REVIEW_REQUIRED
        job["validation_errors"] = [
            issue.to_dict() for issue in content.validation.issues
        ]
        job["status_message"] = "；".join(content.reasons)
        storage.save_job(job)
        raise HTTPException(
            status_code=409,
            detail={
                "message": job["status_message"],
                "unresolved_fields": content.unresolved_fields[:20],
            },
        )

    from .knowledge.history import KnowledgeHistory

    state = ReviewStateStore(job_dir)
    history = KnowledgeHistory(
        Path(os.environ.get("KNOWLEDGE_DB_PATH", DATA_ROOT / "knowledge.sqlite3"))
    )
    try:
        receipt = history.append_confirmed_job(
            job,
            final,
            state.confirmed_hashes(final),
        )
    except (ValueError, FinalResultError) as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    finally:
        history.close()

    write_json_atomic(job_dir / "review" / "finalization.json", receipt)
    gate = evaluate_ready_gate(job, final, job_dir)
    if not gate.ready:
        job["status"] = JobStatus.REVIEW_REQUIRED
        job["status_message"] = "；".join(gate.reasons)
        storage.save_job(job)
        raise HTTPException(status_code=409, detail=job["status_message"])

    job["status"] = JobStatus.READY
    job["status_message"] = "配方已确认并加入知识库，可以导出 Excel。"
    storage.save_job(job)
    return {
        "status": JobStatus.READY,
        "saved": True,
        "message": job["status_message"],
        "receipt": receipt,
    }


@app.post("/api/jobs/{job_id}/finalize")
async def finalize_job(job_id: str):
    return _finalize_job(job_id)


@app.post("/api/jobs/{job_id}/confirm")
async def confirm_job(job_id: str):
    """Compatibility alias for whole-job finalization."""
    return _finalize_job(job_id)


@app.post("/api/jobs/{job_id}/fields/{field_id}/recheck")
async def recheck_field(job_id: str, field_id: str):
    """对单个字段执行完整 OCR/VLM 局部复核。"""
    return await _perform_recheck(job_id, [field_id], _get_storage())


# ─── 分组 API ─────────────────────────────────────────────


@app.get("/api/jobs/{job_id}/tree")
async def get_job_tree(job_id: str):
    """获取任务业务树（公司→产品→配方）。"""
    from .grouping.storage import build_tree_response, load_business_entities
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    job_dir = storage.get_job_dir(job_id)
    entities = load_business_entities(job_dir)
    if not entities:
        return {"summary": {}, "companies": []}
    return build_tree_response(entities)


@app.get("/api/jobs/{job_id}/pages")
async def get_job_pages(job_id: str):
    """获取任务页面视图。"""
    from .grouping.storage import build_pages_response, load_business_entities
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    job_dir = storage.get_job_dir(job_id)
    entities = load_business_entities(job_dir)
    if not entities:
        return []
    return build_pages_response(entities)


@app.get("/api/jobs/{job_id}/formulas/{formula_id}")
async def get_formula_detail(job_id: str, formula_id: str):
    """获取配方详情。"""
    from .grouping.storage import build_formula_detail, load_business_entities
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    job_dir = storage.get_job_dir(job_id)
    entities = load_business_entities(job_dir)
    if not entities:
        raise HTTPException(status_code=404, detail="业务实体不存在。")

    detail = build_formula_detail(entities, formula_id)
    if not detail:
        raise HTTPException(status_code=404, detail=f"配方不存在: {formula_id}")
    return detail


@app.patch("/api/jobs/{job_id}/pages/{page_id}/company")
async def update_page_company(job_id: str, page_id: str, request: Request):
    """更新页面公司归属。"""
    from .grouping.review import GroupingReviewError, GroupingReviewService
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    job_dir = storage.get_job_dir(job_id)
    new_company = body.get("raw_value", body.get("company_id", ""))
    try:
        GroupingReviewService(job_dir, job_id).update_page_company(
            page_id,
            str(new_company),
            str(body.get("standard_value", new_company)),
        )
    except (FinalResultError, GroupingReviewError) as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"status": "updated", "page_id": page_id, "company": new_company}


@app.patch("/api/jobs/{job_id}/formulas/{formula_id}/group")
async def update_formula_group(job_id: str, formula_id: str, request: Request):
    """更新配方分组（公司/产品归属）。"""
    from .grouping.review import GroupingReviewError, GroupingReviewService
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    job_dir = storage.get_job_dir(job_id)
    try:
        GroupingReviewService(job_dir, job_id).update_formula_group(
            formula_id,
            body.get("company_id"),
            body.get("product_id"),
        )
    except (FinalResultError, GroupingReviewError) as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"status": "updated", "formula_id": formula_id}


@app.patch("/api/jobs/{job_id}/formulas/{formula_id}/number")
async def update_formula_number(job_id: str, formula_id: str, request: Request):
    """更新配方编号（不改变 formula_id）。"""
    from .grouping.review import GroupingReviewError, GroupingReviewService
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    job_dir = storage.get_job_dir(job_id)
    new_no = body.get("formula_no_raw", "")
    try:
        GroupingReviewService(job_dir, job_id).update_formula_number(
            formula_id, str(new_no)
        )
    except (FinalResultError, GroupingReviewError) as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"status": "updated", "formula_id": formula_id, "formula_no_raw": new_no}


@app.post("/api/jobs/{job_id}/formulas/merge")
async def merge_formulas(job_id: str, request: Request):
    """合并相邻配方（仅同页面）。"""
    from .grouping.review import GroupingReviewError, GroupingReviewService
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    formula_ids = body.get("formula_ids", [])
    if len(formula_ids) < 2:
        raise HTTPException(status_code=400, detail="至少需要两条配方才能合并。")

    job_dir = storage.get_job_dir(job_id)
    try:
        merged_id = GroupingReviewService(job_dir, job_id).merge_formulas(
            formula_ids
        )
    except (FinalResultError, GroupingReviewError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {
        "status": "merged",
        "formula_id": merged_id,
        "merged_count": len(formula_ids),
    }


@app.post("/api/jobs/{job_id}/formulas/{formula_id}/split")
async def split_formula(job_id: str, formula_id: str, request: Request):
    """拆分配方（按材料索引拆分）。"""
    from .grouping.review import GroupingReviewError, GroupingReviewService
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    split_after = body.get("split_after_material", 0)  # 在第 N 条材料后拆分

    job_dir = storage.get_job_dir(job_id)
    try:
        new_id = GroupingReviewService(job_dir, job_id).split_formula(
            formula_id, int(split_after)
        )
    except (FinalResultError, GroupingReviewError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "split", "original_id": formula_id, "new_id": new_id}


@app.post("/api/jobs/{job_id}/companies/merge")
async def merge_company_groups(job_id: str, request: Request):
    """合并两个公司组。"""
    from .grouping.review import GroupingReviewService
    from .grouping.storage import load_business_entities
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    source_id = body.get("source_company_id", "")
    target_id = body.get("target_company_id", "")
    if not source_id or not target_id:
        raise HTTPException(status_code=400, detail="需要 source_company_id 和 target_company_id。")

    job_dir = storage.get_job_dir(job_id)
    entities = load_business_entities(job_dir)
    if not entities:
        raise HTTPException(status_code=400, detail="业务实体不存在。")

    source = next((c for c in entities.company_groups if c.company_id == source_id), None)
    target = next((c for c in entities.company_groups if c.company_id == target_id), None)
    if not source or not target:
        raise HTTPException(status_code=404, detail="公司组不存在。")

    GroupingReviewService(job_dir, job_id).merge_company(source_id, target_id)
    return {"status": "merged", "target_company_id": target_id}


@app.post("/api/jobs/{job_id}/companies/{company_id}/alias")
async def save_company_alias(job_id: str, company_id: str, request: Request):
    """将原文保存为公司别名到知识库。"""
    from .knowledge.database import KnowledgeDB
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    alias = body.get("alias", "").strip()
    if not alias:
        raise HTTPException(status_code=400, detail="别名不能为空。")

    standard_name = body.get("standard_name", alias).strip()
    db = KnowledgeDB(DATA_ROOT / "knowledge.sqlite3")
    db.initialize()

    # 查找或创建公司（使用 customers 表）
    conn = db._get_conn()
    conn.execute(
        "INSERT OR IGNORE INTO customers (name, usage_count) VALUES (?, 0)",
        (standard_name,),
    )
    conn.commit()
    row = conn.execute("SELECT id FROM customers WHERE name = ?", (standard_name,)).fetchone()
    company_db_id = row["id"] if row else None

    if company_db_id:
        db.add_company_alias(company_db_id, alias, alias_type="user_confirmed", source="web")

    return {"status": "saved", "alias": alias, "standard_name": standard_name}


@app.post("/api/jobs/{job_id}/products/merge")
async def merge_product_groups(job_id: str, request: Request):
    """合并两个产品组。"""
    from .grouping.review import GroupingReviewService
    from .grouping.storage import load_business_entities
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    source_id = body.get("source_product_id", "")
    target_id = body.get("target_product_id", "")
    if not source_id or not target_id:
        raise HTTPException(status_code=400, detail="需要 source_product_id 和 target_product_id。")

    job_dir = storage.get_job_dir(job_id)
    entities = load_business_entities(job_dir)
    if not entities:
        raise HTTPException(status_code=400, detail="业务实体不存在。")

    source = next((p for p in entities.product_groups if p.product_id == source_id), None)
    target = next((p for p in entities.product_groups if p.product_id == target_id), None)
    if not source or not target:
        raise HTTPException(status_code=404, detail="产品组不存在。")

    GroupingReviewService(job_dir, job_id).merge_product(source_id, target_id)
    return {"status": "merged", "target_product_id": target_id}


@app.post("/api/jobs/{job_id}/products/{product_id}/alias")
async def save_product_alias(job_id: str, product_id: str, request: Request):
    """将原文保存为产品别名到知识库。"""
    from .knowledge.database import KnowledgeDB
    storage = _get_storage()
    try:
        storage.get_job(job_id)
    except (FileNotFoundError, ValueError):
        raise HTTPException(status_code=404, detail="任务不存在。")

    body = await request.json()
    alias = body.get("alias", "").strip()
    if not alias:
        raise HTTPException(status_code=400, detail="别名不能为空。")

    standard_name = body.get("standard_name", alias).strip()
    db = KnowledgeDB(DATA_ROOT / "knowledge.sqlite3")
    db.initialize()

    # 查找或创建产品（使用 products 表）
    conn = db._get_conn()
    conn.execute(
        "INSERT OR IGNORE INTO products (name, usage_count) VALUES (?, 0)",
        (standard_name,),
    )
    conn.commit()
    row = conn.execute("SELECT id FROM products WHERE name = ?", (standard_name,)).fetchone()
    product_db_id = row["id"] if row else None

    if product_db_id:
        db.add_product_alias(product_db_id, alias, alias_type="user_confirmed", source="web")

    return {"status": "saved", "alias": alias, "standard_name": standard_name}


# ─── 知识库管理 ─────────────────────────────────────────────


def _knowledge_db_path() -> Path:
    return Path(
        os.environ.get(
            "KNOWLEDGE_DB_PATH",
            DATA_ROOT / "knowledge.sqlite3",
        )
    )


@app.get("/api/knowledge/tree")
async def knowledge_tree(q: str = ""):
    from .knowledge.history import KnowledgeHistory

    history = KnowledgeHistory(_knowledge_db_path())
    try:
        return history.tree(q)
    finally:
        history.close()


@app.get("/api/knowledge/formulas/{formula_id}")
async def knowledge_formula(formula_id: int):
    from .knowledge.history import KnowledgeHistory

    history = KnowledgeHistory(_knowledge_db_path())
    try:
        detail = history.formula_detail(formula_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    finally:
        history.close()
    detail["evidence_image_url"] = ""
    source_job_id = detail.get("source_job_id")
    image_index = int(detail.get("source_image_index", 0))
    if source_job_id and image_index > 0:
        try:
            source_job = _get_storage().get_job(str(source_job_id))
            source = str(source_job.get("images", [])[image_index - 1].get("source", ""))
            if source:
                detail["evidence_image_url"] = (
                    f"/jobs/{quote(str(source_job_id), safe='')}/files/{quote(source, safe='/')}"
                )
        except (FileNotFoundError, ValueError, IndexError):
            pass
    return detail


@app.get("/api/knowledge/compare")
async def compare_knowledge(left: int, right: int):
    from .knowledge.history import KnowledgeHistory

    history = KnowledgeHistory(_knowledge_db_path())
    try:
        return history.compare(left, right)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    finally:
        history.close()


@app.get("/api/knowledge/export")
async def export_knowledge():
    from .knowledge.exporter import export_knowledge_history

    output = (
        _knowledge_db_path().parent
        / "exports"
        / f"formula-knowledge-{int(time.time())}.xlsx"
    )
    export_knowledge_history(_knowledge_db_path(), output)
    return FileResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="配方知识库.xlsx",
    )


@app.get("/knowledge", response_class=HTMLResponse)
async def knowledge_page(request: Request):
    """知识库管理页面。"""
    from .knowledge.database import KnowledgeDB
    db = KnowledgeDB(_knowledge_db_path())
    db.initialize()
    materials = db.get_all_materials()
    db.close()
    return templates.TemplateResponse(
        request, "knowledge.html", {"materials": materials}
    )


@app.post("/knowledge/materials")
async def add_material(request: Request):
    """添加物料。"""
    from .knowledge.database import KnowledgeDB
    body = await request.json()
    name = body.get("name", "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="物料名称不能为空。")

    db = KnowledgeDB(_knowledge_db_path())
    db.initialize()
    mid = db.add_material(name, body.get("category", ""), body.get("unit", ""))

    # 添加别名
    for alias in body.get("aliases", []):
        if alias.strip():
            db.add_alias(mid, alias.strip(), alias_type="manual", source="web")

    return {"id": mid, "name": name}


@app.post("/knowledge/import")
async def import_knowledge(file: UploadFile = File(...)):
    """从 Excel/CSV 导入物料和配方到知识库。"""
    import csv
    import io as _io

    from .knowledge.database import KnowledgeDB

    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名不能为空。")

    ext = Path(file.filename).suffix.lower()
    if ext not in (".xlsx", ".csv"):
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {ext}。允许: .xlsx, .csv")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="文件为空。")

    db = KnowledgeDB(_knowledge_db_path())
    db.initialize()

    imported_count = 0

    try:
        if ext == ".csv":
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(_io.StringIO(text))
            for row in reader:
                name = (row.get("物料名称") or row.get("name") or row.get("material") or "").strip()
                if name:
                    unit = (row.get("单位") or row.get("unit") or "").strip()
                    category = (row.get("分类") or row.get("category") or "").strip()
                    db.add_material(name, category, unit)
                    # 别名
                    alias_str = (row.get("别名") or row.get("aliases") or "").strip()
                    if alias_str:
                        mid = db.add_material(name, category, unit)
                        for alias in alias_str.replace("；", ";").split(";"):
                            if alias.strip():
                                db.add_alias(mid, alias.strip(), alias_type="import", source="csv")
                    imported_count += 1
        else:
            # xlsx
            import io as _io2

            from openpyxl import load_workbook
            wb = load_workbook(_io2.BytesIO(content), read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                raise HTTPException(status_code=400, detail="文件行数不足。")

            # 自动识别列
            headers = [str(h or "").strip() for h in rows[0]]
            name_col = None
            unit_col = None
            category_col = None
            alias_col = None
            for i, h in enumerate(headers):
                hl = h.lower()
                if "物料" in h or "名称" in h or hl == "name" or hl == "material":
                    name_col = i
                elif "单位" in h or hl == "unit":
                    unit_col = i
                elif "分类" in h or hl == "category":
                    category_col = i
                elif "别名" in h or hl == "aliases":
                    alias_col = i

            if name_col is None:
                name_col = 0  # 默认第一列

            for row in rows[1:]:
                if not row or len(row) <= name_col:
                    continue
                name = str(row[name_col] or "").strip()
                if not name:
                    continue
                unit = str(row[unit_col] or "").strip() if unit_col is not None and len(row) > unit_col else ""
                category = str(row[category_col] or "").strip() if category_col is not None and len(row) > category_col else ""
                mid = db.add_material(name, category, unit)
                if alias_col is not None and len(row) > alias_col:
                    alias_str = str(row[alias_col] or "").strip()
                    if alias_str:
                        for alias in alias_str.replace("；", ";").split(";"):
                            if alias.strip():
                                db.add_alias(mid, alias.strip(), alias_type="import", source="xlsx")
                imported_count += 1
            wb.close()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"导入失败: {e}")

    return {"imported": imported_count, "message": f"成功导入 {imported_count} 条物料记录。"}


# ─── 启动入口 ───────────────────────────────────────────────


def run() -> None:
    """启动应用。"""
    import os

    import uvicorn

    cfg = get_config()
    reload = os.environ.get("APP_RELOAD", "").lower() in ("1", "true", "yes")
    uvicorn.run(
        "lite_app.main:app",
        host=cfg.host,
        port=cfg.port,
        reload=reload,
    )


if __name__ == "__main__":
    run()
